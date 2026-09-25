from __future__ import annotations

import glob
import io
import json
import os
import subprocess
import sys
from base64 import b64encode
from collections import Counter
from collections.abc import Iterator
from copy import deepcopy
from decimal import Decimal
from math import nextafter
from pathlib import Path, PureWindowsPath
from textwrap import dedent
from types import SimpleNamespace
from typing import Any, Literal, cast

import polars as pl
import pytest

import openwrangler_runtime.engines.base as engine_base
import openwrangler_runtime.engines.polars_engine as polars_engine
from openwrangler_runtime._column_binding import bind_step
from openwrangler_runtime.engines.base import EngineError, typed_selection_value
from openwrangler_runtime.engines.polars_engine import PolarsEngine
from openwrangler_runtime.export_target import ExportWriterPath, _regular_file_identity
from openwrangler_runtime.lineage import source_lineage
from openwrangler_runtime.operations import OperationError, validate_step
from openwrangler_runtime.session import SessionManager
from openwrangler_runtime.session_source import SourceChangedError

ROOT = Path(__file__).resolve().parents[2]


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_explode_list_preserves_current_children_names_and_nulls(lazy: bool) -> None:
    operation = {
        "id": "explode",
        "kind": "explodeList",
        "params": {"column": {"id": "c:source:1", "name": "*", "position": 1}},
    }
    engine = PolarsEngine()
    try:
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([operation]), namespace)
        record = {"^a.*$": -1, "*": 7}
        cases = [
            (pl.Int64, [[-(2**63), None], [], None, [2**63 - 1, 3]], [-(2**63), None, None, None, 2**63 - 1, 3]),
            (pl.Datetime("ns", "America/New_York"), [[-1, None], [], None, [1, 1001]], [-1, None, None, None, 1, 1001]),
            (
                pl.Decimal(20, 3),
                [[Decimal("-12345678901234567.123"), None], [], None, [Decimal("1.001"), Decimal("0.000")]],
                [Decimal("-12345678901234567.123"), None, None, None, Decimal("1.001"), Decimal("0.000")],
            ),
            (
                pl.Struct({"^a.*$": pl.Datetime("ns"), "*": pl.Int64}),
                [[record, None], [], None, [record, record]],
                [record, None, None, None, record, record],
            ),
            (pl.List(pl.Int64), [[[1, None], []], [], None, [[2], [3, 4]]], [[1, None], [], None, None, [2], [3, 4]]),
            (
                pl.Array(pl.Int64, 2),
                [[[1, None], [2, 3]], [], None, [[4, 5], [6, 7]]],
                [[1, None], [2, 3], None, None, [4, 5], [6, 7]],
            ),
            (pl.Null, [[None, None], [], None, [None, None]], [None] * 6),
        ]
        for dtype, values, expected_values in cases:
            source = pl.DataFrame({"keep": [9, 8, 7, 6], "*": pl.Series("*", values).cast(pl.List(dtype))})
            before = source.clone()
            expected = pl.DataFrame({"keep": [9, 9, 8, 7, 6, 6], "*": pl.Series("*", expected_values).cast(dtype)})
            for frame, wanted in (
                (source, expected),
                (source.head(0), expected.head(0)),
                (source.slice(2, 1), expected.slice(3, 1)),
            ):
                native = frame.lazy() if lazy else frame
                for result in (engine.apply_transform(native, operation), namespace["clean_data"](native)):
                    assert isinstance(result, pl.LazyFrame) == lazy
                    actual = result.collect() if lazy else result
                    assert actual.schema == wanted.schema and actual.equals(wanted)
            reordered = pl.DataFrame([source.get_column("*"), source.get_column("keep")])
            result = namespace["clean_data"](reordered.lazy() if lazy else reordered)
            actual = result.collect() if lazy else result
            assert actual.equals(pl.DataFrame([expected.get_column("*"), expected.get_column("keep")]))
            assert source.schema == before.schema and source.equals(before)
        # Reused programs resolve the current child type, without removing caller columns.
        private_looking = engine_base.INTERNAL_ROW_ID_PREFIX + "caller"
        changed = pl.DataFrame({"*": [["one", "two"], None], private_looking: [4, 5]})
        result = namespace["clean_data"](changed.lazy() if lazy else changed)
        actual = result.collect() if lazy else result
        assert actual.equals(pl.DataFrame({"*": ["one", "two", None], private_looking: [4, 4, 5]}))
    finally:
        engine.close()


def test_polars_explode_list_rejects_native_domain_before_collection(monkeypatch: pytest.MonkeyPatch) -> None:
    operation = {
        "id": "explode",
        "kind": "explodeList",
        "params": {"column": {"id": "c:source:0", "name": "values", "position": 0}},
    }
    engine = PolarsEngine()
    try:
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([operation]), namespace)

        def forbidden(*_args, **_kwargs):
            raise AssertionError("Explode List admission collected native rows")

        safe = pl.DataFrame({"values": [[1]]}).lazy()
        # Metadata injection only: never allocate Object-backed native containers.
        for dtype in (
            pl.Int64,
            pl.Array(pl.Int64, 1),
            pl.Object,
            pl.List(pl.Object),
            pl.List(pl.Struct({"safe": pl.Int64, "bad": pl.Array(pl.List(pl.Object), 2)})),
        ):
            with monkeypatch.context() as patch:
                patch.setattr(pl.LazyFrame, "collect_schema", lambda _self, dtype=dtype: pl.Schema({"values": dtype}))
                patch.setattr(pl.LazyFrame, "collect", forbidden)
                with pytest.raises(EngineError, match="Explode List"):
                    engine.apply_transform(safe, operation)
                with pytest.raises(ValueError, match="Explode List"):
                    namespace["clean_data"](safe)
        with monkeypatch.context() as patch:
            patch.setattr(pl.LazyFrame, "collect", forbidden)
            with pytest.raises(ValueError, match="Explode List"):
                namespace["clean_data"](pl.DataFrame({"other": [1]}).lazy())
    finally:
        engine.close()


def test_polars_explode_list_guards_retained_input_before_expansion(monkeypatch: pytest.MonkeyPatch) -> None:
    operation = {
        "id": "explode",
        "kind": "explodeList",
        "params": {"column": {"id": "c:source:1", "name": "values", "position": 1}},
    }
    monkeypatch.setattr(polars_engine, "_POLARS_EXPLODE_MAX_ROWS", 8, raising=False)
    engine = PolarsEngine()
    try:
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([operation]), namespace)
        late = pl.DataFrame({"keep": [3, 2, 1], "values": [[1], [2], list(range(9))]})

        def forbidden(*_args, **_kwargs):
            raise AssertionError("Explode List constructed expansion before rejecting growth")

        counts = []
        original_select = pl.DataFrame.select

        def observe_count(frame, *args, **kwargs):
            result = original_select(frame, *args, **kwargs)
            counts.append((result.dtypes, result.item()))
            return result

        with monkeypatch.context() as patch:
            patch.setattr(pl.DataFrame, "select", observe_count)
            patch.setattr(pl.DataFrame, "explode", forbidden)
            patch.setattr(pl.LazyFrame, "explode", forbidden)
            for frame in (late, late.lazy()):
                with pytest.raises(EngineError, match="8.*row"):
                    engine.apply_transform(frame, operation)
                with pytest.raises(ValueError, match="8.*row"):
                    namespace["clean_data"](frame)
        assert counts == [([pl.UInt64], 11)] * 4
        at_capacity = pl.DataFrame({"keep": [1], "values": [list(range(8))]})
        assert engine.apply_transform(at_capacity, operation).height == 8
        assert namespace["clean_data"](at_capacity).equals(engine.apply_transform(at_capacity, operation))
        base = pl.DataFrame({"keep": [9, 7], "values": [[0], [0]]})
        for transform in (lambda value: engine.apply_transform(value, operation), namespace["clean_data"]):
            calls = []

            def variable(batch, calls=calls):
                calls.append(1)
                values = [[30, None], [10]] if len(calls) == 1 else [[99], list(range(9))]
                return batch.with_columns(pl.Series("values", values, dtype=pl.List(pl.Int64)))

            source = base.lazy().map_batches(
                variable,
                schema=base.schema,
                predicate_pushdown=False,
                projection_pushdown=False,
                slice_pushdown=False,
                streamable=False,
            )
            result = transform(source)
            assert isinstance(result, pl.LazyFrame) and calls == [1]
            assert engine.shape(result) == {"rows": 3, "columns": 2}
            assert [[cell["raw"] for cell in row["values"]] for row in engine.page(result, 1, 2)["rows"]] == [
                [9, None],
                [7, 10],
            ]
            assert result.collect().equals(pl.DataFrame({"keep": [9, 9, 7], "values": [30, None, 10]}))
            assert result.select("keep").collect().get_column("keep").to_list() == [9, 9, 7]
            assert calls == [1]
        assert base.get_column("values").to_list() == [[0], [0]]
    finally:
        engine.close()


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_extract_struct_fields_preserves_literal_names_and_current_input(lazy: bool) -> None:
    source = pl.DataFrame(
        {"id": [0, 1, 2], "^a.*$": [{"*": 7, "^a.*$": 8, "amount": 90}, None, {"*": None, "^a.*$": 4, "amount": 91}]}
    )
    engine = PolarsEngine()
    try:
        schema = engine.schema(source)
        lineage = source_lineage(schema)
        operation = bind_step(
            validate_step(
                {
                    "id": "extract",
                    "kind": "extractStructFields",
                    "params": {
                        "column": lineage[1],
                        "fields": [{"field": "^a.*$", "newColumn": "*"}, {"field": "*", "newColumn": 'selected"value'}],
                    },
                }
            ),
            schema,
            lineage,
        )
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([operation]), namespace)
        reordered_dtype = pl.Struct({"amount": pl.Int64, "^a.*$": pl.Int64, "*": pl.Int64})
        reordered_parent = source.get_column("^a.*$").cast(reordered_dtype)
        assert reordered_parent.dtype == reordered_dtype
        reordered = pl.DataFrame([reordered_parent, source.get_column("id")])
        for frame in (source, reordered, source.head(0)):
            native = frame.lazy() if lazy else frame
            for result in (engine.apply_transform(native, operation), namespace["clean_data"](native)):
                assert isinstance(result, pl.LazyFrame) == lazy
                result = result.collect() if lazy else result
                assert result.columns == [*frame.columns, "*", 'selected"value']
                assert result.get_column("*").to_list() == [8, None, 4][: frame.height]
                assert result.get_column('selected"value').to_list() == [7, None, None][: frame.height]
                assert all(result.get_column(name).equals(frame.get_column(name)) for name in frame.columns)
        assert source.get_column("id").to_list() == [0, 1, 2]
    finally:
        engine.close()


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_extract_struct_fields_retains_native_scalar_storage(lazy: bool) -> None:
    children = pl.DataFrame(
        [
            pl.Series("text", ["é", "hidden", None], dtype=pl.String),
            pl.Series("category", ["one", "two", None], dtype=pl.Categorical),
            pl.Series("enum", ["one", "two", None], dtype=pl.Enum(["one", "two"])),
            pl.Series("signed", [-(2**127), 2**127 - 1, None], dtype=pl.Int128),
            pl.Series("unsigned", [2**128 - 1, 0, None], dtype=pl.UInt128),
            pl.Series("float", [float("inf"), 1.5, None], dtype=pl.Float64),
            pl.Series(
                "decimal",
                [Decimal("12345678901234567890123456.7890"), Decimal("-1.0000"), None],
                dtype=pl.Decimal(30, 4),
            ),
            pl.Series("boolean", [True, False, None], dtype=pl.Boolean),
            pl.Series("date", [-1, 1, None], dtype=pl.Int32).cast(pl.Date),
            *[
                pl.Series(f"timestamp_{unit}", [-1, 1, None], dtype=pl.Int64).cast(pl.Datetime(unit))
                for unit in ("ns", "us", "ms")
            ],
            pl.Series("zoned", [-1, 1, None], dtype=pl.Int64).cast(pl.Datetime("ns", "America/New_York")),
            *[
                pl.Series(f"duration_{unit}", [-1, 1, None], dtype=pl.Int64).cast(pl.Duration(unit))
                for unit in ("ns", "us", "ms")
            ],
            pl.Series("binary", [b"\x00\xff", b"hidden", None], dtype=pl.Binary),
        ]
    )
    source = children.select(pl.struct(pl.all()).alias("record")).with_row_index("id")
    source = source.with_columns(pl.when(pl.col("id") == 1).then(None).otherwise(pl.col("record")).alias("record"))
    before = source.clone()
    fields = [{"field": name, "newColumn": f"out_{index}"} for index, name in enumerate(children.columns)]
    operation = {
        "id": "extract",
        "kind": "extractStructFields",
        "params": {"column": {"id": "c:source:1", "name": "record", "position": 1}, "fields": fields},
    }
    engine = PolarsEngine()
    try:
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([operation]), namespace)
        for frame in (source, source.head(0), source.filter(pl.col("id") == 1)):
            native = frame.lazy() if lazy else frame
            for result in (engine.apply_transform(native, operation), namespace["clean_data"](native)):
                result = result.collect() if lazy else result
                assert result.columns == [*source.columns, *[field["newColumn"] for field in fields]]
                assert result.get_column("record").equals(frame.get_column("record"))
                for field in fields:
                    actual = result.get_column(field["newColumn"])
                    child = children.get_column(field["field"])
                    assert actual.dtype == child.dtype
                    expected = pl.concat(
                        [child.head(1), pl.Series(child.name, [None, None], dtype=child.dtype)]
                    ).gather(frame.get_column("id"))
                    assert actual.equals(expected.rename(actual.name))
        assert source.equals(before)
    finally:
        engine.close()


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_extract_struct_fields_revalidates_before_evaluating(lazy: bool) -> None:
    good = pl.DataFrame({"record": [{"value": 1}], "keep": [2]})
    operation = {
        "id": "extract",
        "kind": "extractStructFields",
        "params": {
            "column": {"id": "c:source:0", "name": "record", "position": 0},
            "fields": [{"field": "value", "newColumn": "selected"}],
        },
    }
    engine = PolarsEngine()
    try:
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([operation]), namespace)
        invalid = [
            good.drop("record"),
            good.with_columns(pl.lit(1).alias("record")),
            pl.DataFrame({"record": pl.Series([object()], dtype=pl.Object)}),
            good.with_columns(pl.struct(pl.lit(1).alias("other")).alias("record")),
            good.with_columns(pl.lit(9).alias("selected")),
            *[
                pl.DataFrame({"record": pl.Series([{"value": None}], dtype=pl.Struct({"value": dtype}))})
                for dtype in (
                    pl.List(pl.Int64),
                    pl.Array(pl.Int64, 2),
                    pl.Struct({"child": pl.Int64}),
                    pl.Time,
                    pl.Null,
                )
            ],
        ]

        def forbidden(batch):
            raise AssertionError("Metadata admission evaluated source rows")

        for frame in invalid:
            native = frame.lazy().map_batches(forbidden, schema=frame.schema) if lazy else frame
            with pytest.raises(EngineError, match="Extract Struct Fields"):
                engine.apply_transform(native, operation)
            with pytest.raises(ValueError, match="Extract Struct Fields"):
                namespace["clean_data"](native)
        # Re-resolve the current native field dtype; changing one admitted scalar to another is valid.
        changed = good.with_columns(pl.struct(pl.lit("new").alias("value")).alias("record"))
        result = namespace["clean_data"](changed.lazy() if lazy else changed)
        result = result.collect() if lazy else result
        assert result["selected"].to_list() == ["new"]
        assert result.schema["selected"] == pl.String
        if lazy:
            batches = []

            def observed(batch):
                batches.append(batch.height)
                return batch

            watched = good.lazy().map_batches(observed, schema=good.schema)
            live = engine.apply_transform(watched, operation)
            generated = namespace["clean_data"](watched)
            assert batches == []
            assert live.collect()["selected"].to_list() == [1]
            assert generated.collect()["selected"].to_list() == [1]
            assert batches == [1, 1]
    finally:
        engine.close()


def test_polars_extract_struct_fields_masks_hidden_parquet_children(tmp_path: Path) -> None:
    import pyarrow as pa
    import pyarrow.parquet as pq

    child = pa.array([-1, 123, None], type=pa.int64()).cast(pa.timestamp("ns"))
    parent = pa.StructArray.from_arrays(
        [child], names=["__open_wrangler_internal_row_id_child"], mask=pa.array([False, True, False])
    )
    assert parent.field(0).cast(pa.int64()).to_pylist() == [-1, 123, None]
    assert parent[1].as_py() is None
    path = tmp_path / "hidden.parquet"
    pq.write_table(pa.table({"record": parent}), path)
    before = path.read_bytes()
    source = pl.scan_parquet(path)
    operation = {
        "id": "extract",
        "kind": "extractStructFields",
        "params": {
            "column": {"id": "c:source:0", "name": "record", "position": 0},
            "fields": [{"field": "__open_wrangler_internal_row_id_child", "newColumn": "selected"}],
        },
    }
    engine = PolarsEngine()
    try:
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([operation]), namespace)
        for result in (engine.apply_transform(source, operation), namespace["clean_data"](source)):
            result = result.collect()
            assert result["selected"].cast(pl.Int64).to_list() == [-1, None, None]
            assert result["record"].equals(source.collect()["record"])
        assert path.read_bytes() == before
    finally:
        engine.close()


def test_polars_extract_struct_fields_appends_64_columns_and_preserves_row_identity() -> None:
    source = pl.DataFrame(
        {"record": [{f"field_{index}": index for index in range(65)}], "__open_wrangler_internal_row_id_test": [123]}
    )
    fields = [{"field": f"field_{index}", "newColumn": f"selected_{index}"} for index in range(64)]
    engine = PolarsEngine()
    try:
        operation = {
            "id": "extract",
            "kind": "extractStructFields",
            "params": {"column": {"id": "c:source:0", "name": "record", "position": 0}, "fields": fields},
        }
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([operation]), namespace)
        for result in (engine.apply_transform(source, operation), namespace["clean_data"](source)):
            assert result.columns == [*source.columns, *[field["newColumn"] for field in fields]]
            assert result.row(0)[2:] == tuple(range(64))
            assert result["__open_wrangler_internal_row_id_test"].to_list() == [123]
        assert source.width == 2
    finally:
        engine.close()


def _literal_polars_source() -> pl.DataFrame:
    return pl.DataFrame({"^a.*$": [1, 1, 2, None], "amount": [20, 30, 40, 50], "*": [3, 1, 2, None]})


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize("generated", [False, True])
@pytest.mark.parametrize("name,expected", [("^a.*$", [1, 1, 2, None]), ("*", [3, 1, 2, None])])
def test_polars_literal_column_names_clone_the_selected_values(
    lazy: bool, generated: bool, name: str, expected: list[int | None]
) -> None:
    source = _literal_polars_source()
    before = source.clone()
    engine = PolarsEngine()
    try:
        schema = engine.schema(source.lazy())
        lineage = source_lineage(schema)
        reference = next(column for column in lineage if column["name"] == name)
        operation = bind_step(
            validate_step({"id": "copy", "kind": "cloneColumn", "params": {"column": reference, "newName": "copy"}}),
            schema,
            lineage,
        )
        frame = source.lazy() if lazy else source
        namespace: dict[str, Any] = {}
        if generated:
            exec(engine.compile_plan([operation]), namespace)
            result = namespace["clean_data"](frame)
        else:
            result = engine.apply_transform(frame, operation)
        assert isinstance(result, pl.LazyFrame) == lazy
        eager = result.collect() if lazy else result
        assert eager.get_column("copy").to_list() == expected
        assert eager.schema["copy"] == pl.Int64
        assert eager.columns == [*source.columns, "copy"]
        assert all(eager.get_column(column).equals(source.get_column(column)) for column in source.columns)
        if generated:
            reordered = pl.DataFrame({column: source.get_column(column) for column in reversed(source.columns)})
            reused = namespace["clean_data"](reordered.lazy() if lazy else reordered)
            reused = reused.collect() if lazy else reused
            assert reused.get_column("copy").to_list() == expected
            assert reused.columns == [*reordered.columns, "copy"]
        assert source.equals(before)
    finally:
        engine.close()


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize("name", ["^a.*$", "*"])
def test_polars_literal_column_names_refuse_missing_generated_input(lazy: bool, name: str) -> None:
    source = _literal_polars_source()
    engine = PolarsEngine()
    try:
        schema = engine.schema(source.lazy())
        lineage = source_lineage(schema)
        operation = bind_step(
            validate_step(
                {
                    "id": "copy",
                    "kind": "cloneColumn",
                    "params": {
                        "column": next(column for column in lineage if column["name"] == name),
                        "newName": "copy",
                    },
                }
            ),
            schema,
            lineage,
        )
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([operation]), namespace)
        missing = pl.DataFrame({column: source.get_column(column) for column in source.columns if column != name})
        before = missing.clone()
        with pytest.raises((ValueError, pl.exceptions.PolarsError)):
            result = namespace["clean_data"](missing.lazy() if lazy else missing)
            if isinstance(result, pl.LazyFrame):
                result.collect_schema()
        assert missing.equals(before)
    finally:
        engine.close()


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_literal_column_names_preserve_projected_views_and_profiles(lazy: bool) -> None:
    source = _literal_polars_source()
    before = source.clone()
    engine = PolarsEngine()
    try:
        frame = engine.ensure_row_ids(source.lazy() if lazy else source, "literal-view")
        schema = engine.schema(frame)
        assert [column["name"] for column in schema] == source.columns
        assert [column["position"] for column in schema] == [0, 1, 2]
        page = engine.page(frame, 1, 2, total_rows=4, column_projection=[(2, "star"), (0, "pattern")])
        assert page["columnIds"] == ["star", "pattern"]
        assert [[cell["display"] for cell in row["values"]] for row in page["rows"]] == [["1", "1"], ["2", "2"]]
        assert [row["rowNumber"] for row in page["rows"]] == [1, 2]
        assert [row["id"].rsplit(":", 1)[-1] for row in page["rows"]] == ["1", "2"]
        summaries = engine.summaries(frame, [(0, "pattern"), (2, "star")])
        assert [summary["numeric"]["exactSum"]["display"] for summary in summaries] == ["4", "6"]
        assert [summary["nullCount"] for summary in summaries] == [1, 1]
        assert {item["value"]: item["count"] for item in summaries[0]["topValues"]} == {"1": 2, "2": 1}
        choices, more = engine.column_values(frame, "^a.*$")
        assert not more and {item["value"]: item["count"] for item in choices} == {"1": 2, "2": 1}
        assert engine.missing_count(frame, 0) == 1
        stats = engine.header_stats(frame)
        assert (stats["missingCells"], stats["missingRows"], stats["duplicateRows"]) == (2, 1, 0)
        filtered = engine.apply_filter_model(
            frame,
            {
                "filters": [
                    {
                        "column": "^a.*$",
                        "type": "integer",
                        "predicates": [{"kind": "predicate", "operator": "gt", "value": "1"}],
                    }
                ],
                "sort": [],
            },
        )
        assert [row["id"].rsplit(":", 1)[-1] for row in engine.page(filtered, 0, 4)["rows"]] == ["2"]
        sorted_frame = engine.apply_filter_model(
            frame, {"filters": [], "sort": [{"column": "*", "direction": "asc", "nulls": "last"}]}
        )
        assert [row["id"].rsplit(":", 1)[-1] for row in engine.page(sorted_frame, 0, 4)["rows"]] == ["1", "2", "0", "3"]
        assert source.equals(before)
    finally:
        engine.close()


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize(
    "kind", ["selectColumns", "dropColumns", "dropDuplicates", "markDuplicates", "dropMissingRows", "oneHotEncode"]
)
def test_polars_literal_column_names_keep_structural_operations_exact(lazy: bool, kind: str) -> None:
    source = _literal_polars_source()
    before = source.clone()
    engine = PolarsEngine()
    try:
        schema = engine.schema(source.lazy())
        lineage = source_lineage(schema)
        params: dict[str, Any] = {"columns": [lineage[0]]}
        expected: list[tuple[Any, ...]]
        if kind == "selectColumns":
            params["columns"] = [lineage[2], lineage[0]]
            expected = [(3, 1), (1, 1), (2, 2), (None, None)]
            expected_columns = ["*", "^a.*$"]
        elif kind == "dropColumns":
            expected = [(20, 3), (30, 1), (40, 2), (50, None)]
            expected_columns = ["amount", "*"]
        elif kind == "dropDuplicates":
            params["keep"] = "first"
            expected = [(1, 20, 3), (2, 40, 2), (None, 50, None)]
            expected_columns = source.columns
        elif kind == "markDuplicates":
            params["newColumn"] = "duplicate"
            expected = [(1, 20, 3, True), (1, 30, 1, True), (2, 40, 2, False), (None, 50, None, False)]
            expected_columns = [*source.columns, "duplicate"]
        elif kind == "oneHotEncode":
            params["dropOriginal"] = True
            expected = [(20, 3, 1, 0), (30, 1, 1, 0), (40, 2, 0, 1), (50, None, 0, 0)]
            expected_columns = ["amount", "*", "^a.*$_1", "^a.*$_2"]
        else:
            params["how"] = "any"
            expected = [(1, 20, 3), (1, 30, 1), (2, 40, 2)]
            expected_columns = source.columns
        operation = bind_step(validate_step({"id": "selected", "kind": kind, "params": params}), schema, lineage)
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([operation]), namespace)
        frame = source.lazy() if lazy else source
        for result in (engine.apply_transform(frame, operation), namespace["clean_data"](frame)):
            assert isinstance(result, pl.LazyFrame) == (lazy and kind != "oneHotEncode")
            eager = result.collect() if isinstance(result, pl.LazyFrame) else result
            assert eager.columns == expected_columns
            assert eager.rows() == expected
            if kind == "oneHotEncode":
                assert [eager.schema[name] for name in expected_columns[-2:]] == [pl.Int8, pl.Int8]
        assert source.equals(before)
    finally:
        engine.close()


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_literal_column_names_preserve_directional_fill_helpers(lazy: bool) -> None:
    source = pl.DataFrame({"^a.*$": [2, 1, 3], "amount": [1, 2, 3], "*": [None, "seed", None]})
    before = source.clone()
    engine = PolarsEngine()
    try:
        schema = engine.schema(source.lazy())
        lineage = source_lineage(schema)
        operation = bind_step(
            validate_step(
                {
                    "id": "fill",
                    "kind": "fillMissingValues",
                    "params": {
                        "column": lineage[2],
                        "replacement": {
                            "kind": "directional",
                            "direction": "forward",
                            "orderBy": [{"column": lineage[0], "direction": "asc", "nulls": "last"}],
                        },
                    },
                }
            ),
            schema,
            lineage,
        )
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([operation]), namespace)
        frame = source.lazy() if lazy else source
        for result in (engine.apply_transform(frame, operation), namespace["clean_data"](frame)):
            assert isinstance(result, pl.LazyFrame) == lazy
            eager = result.collect() if lazy else result
            assert eager.rows() == [(2, 1, "seed"), (1, 2, "seed"), (3, 3, "seed")]
            assert eager.schema == source.schema
        assert source.equals(before)
    finally:
        engine.close()


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize("name", ["^a.*$", "*"])
def test_polars_literal_column_names_keep_temporal_profile_values_exact(lazy: bool, name: str) -> None:
    source = pl.DataFrame(
        {
            name: pl.Series([-1, 0, -1, None], dtype=pl.Int64).cast(pl.Datetime("ns")),
            "amount": pl.Series([1000, 2000, 3000, 4000], dtype=pl.Int64).cast(pl.Datetime("ns")),
        }
    )
    before = source.clone()
    engine = PolarsEngine()
    try:
        frame = source.lazy() if lazy else source
        summary = engine.summaries(frame, [(0, "selected")])[0]
        assert summary["nullCount"] == 1 and summary["distinctCount"] == 2
        expected = {"1969-12-31T23:59:59.999999999": 2, "1970-01-01T00:00:00": 1}
        assert {item["value"]: item["count"] for item in summary["topValues"]} == expected
        choices, more = engine.column_values(frame, name)
        assert not more and {item["value"]: item["count"] for item in choices} == expected
        assert source.equals(before) and source.schema == before.schema
    finally:
        engine.close()


@pytest.mark.parametrize("name", ["^a.*$", "*"])
@pytest.mark.parametrize("layout", [None, "DD/MM/YYYY"])
def test_polars_literal_column_names_preserve_temporal_cast_metadata(name: str, layout: str | None) -> None:
    source = pl.DataFrame(
        {name: ["02/01/2020" if layout else "2020-01-02T00:00:00", None], "amount": ["wrong", "column"]}
    )
    before = source.clone()
    engine = PolarsEngine()
    try:
        schema = engine.schema(source.lazy())
        lineage = source_lineage(schema)
        params = {"column": lineage[0], "dtype": "datetime", **({"inputFormat": layout} if layout else {})}
        operation = bind_step(validate_step({"id": "cast", "kind": "castColumn", "params": params}), schema, lineage)
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([operation]), namespace)
        for frame in (source, source.lazy()):
            for result in (engine.apply_transform(frame, operation), namespace["clean_data"](frame)):
                eager = result.collect() if isinstance(result, pl.LazyFrame) else result
                assert eager.get_column(name).cast(pl.Int64).to_list() == [1577923200000000, None]
                assert eager.schema[name] == pl.Datetime("us")
                assert eager.columns == source.columns
                assert eager.get_column("amount").equals(source.get_column("amount"))
        assert source.equals(before)
    finally:
        engine.close()


def test_polars_literal_column_names_bind_by_example_after_rename() -> None:
    source = pl.DataFrame({"text": ["alpha", "beta", None], "amount": ["wrong", "column", "keep"]})
    before = source.clone()
    engine = PolarsEngine()
    try:
        schema = engine.schema(source.lazy())
        lineage = source_lineage(schema)
        rename = bind_step(
            validate_step(
                {"id": "rename", "kind": "renameColumn", "params": {"column": lineage[0], "newName": "^a.*$"}}
            ),
            schema,
            lineage,
        )
        renamed = engine.apply_transform(source, rename)
        renamed_schema = engine.schema(renamed.lazy())
        renamed_lineage = source_lineage(renamed_schema)
        operation = bind_step(
            validate_step(
                {
                    "id": "example",
                    "kind": "byExample",
                    "params": {
                        "sourceColumns": [renamed_lineage[0]],
                        "newColumn": "label",
                        "examples": [{"inputs": ["alpha"], "output": "ALPHA"}, {"inputs": ["beta"], "output": "BETA"}],
                    },
                }
            ),
            renamed_schema,
            renamed_lineage,
        )
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([rename, operation]), namespace)
        for frame in (source, source.lazy()):
            live = engine.apply_transform(engine.apply_transform(frame, rename), operation)
            for result in (live, namespace["clean_data"](frame)):
                eager = result.collect() if isinstance(result, pl.LazyFrame) else result
                assert eager.columns == ["^a.*$", "amount", "label"]
                assert eager.rows() == [("alpha", "wrong", "ALPHA"), ("beta", "column", "BETA"), (None, "keep", None)]
        assert source.equals(before)
    finally:
        engine.close()


@pytest.mark.parametrize("name", ["^a.*$", "*"])
def test_polars_literal_column_names_refuse_unsafe_uint128_before_evaluation(
    name: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = pl.DataFrame({"value": pl.Series([1, 2], dtype=pl.UInt128), name: pl.Series([2, 3], dtype=pl.UInt128)})
    before = source.clone()
    engine = PolarsEngine()
    try:
        operation = _polars_formula_literal_operation(source.lazy(), "add", None, right_column=True)
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([operation]), namespace)
        monkeypatch.setattr(pl, "__version__", "1.35.2")

        def forbid_evaluation(*args: Any, **kwargs: Any) -> Any:
            raise AssertionError("Unqualified UInt128 must be refused before evaluation")

        monkeypatch.setattr(pl.LazyFrame, "collect", forbid_evaluation)
        monkeypatch.setattr(pl.DataFrame, "with_columns", forbid_evaluation)
        for frame in (source, source.lazy()):
            for run in (
                lambda frame=frame: engine.apply_transform(frame, operation),
                lambda frame=frame: namespace["clean_data"](frame),
            ):
                with pytest.raises((EngineError, ValueError), match="stable Polars 1.36"):
                    run()
        assert source.equals(before)
    finally:
        engine.close()


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize(
    "dtype,tick,raw",
    [
        (pl.Datetime("ns"), -1, "1969-12-31T23:59:59.999999999"),
        (pl.Datetime("ns", "America/New_York"), -1, "1969-12-31T18:59:59.999999999-05:00"),
        (pl.Datetime("us"), 1, "1970-01-01T00:00:00.000001"),
        (pl.Datetime("ms"), 1, "1970-01-01T00:00:00.001000"),
        (pl.Duration("ns"), 1, "0.000000001"),
        (pl.Duration("ns"), -(2**63), "-9223372036.854775808"),
        (pl.Duration("us"), 1, "0.000001"),
        (pl.Duration("ms"), -12345, -12.345),
    ],
)
def test_polars_nested_temporal_output_preserves_values_and_native_counts(
    lazy: bool, dtype: Any, tick: int, raw: Any
) -> None:
    source = pl.DataFrame({"value": pl.Series([tick, tick + 1, tick, None], dtype=pl.Int64).cast(dtype)})
    source = source.select(
        pl.concat_list("value").alias("items"),
        pl.concat_list("value", pl.lit(None, dtype=dtype)).list.to_array(2).alias("fixed"),
        pl.struct(pl.col("value").alias("when"), pl.lit(41).alias("ticks"), pl.lit("keep").alias("text")).alias(
            "record"
        ),
    )
    source = pl.concat(
        [
            source,
            pl.DataFrame(
                {
                    "items": [[], None],
                    "fixed": [None, [None, None]],
                    "record": [None, {"when": None, "ticks": 41, "text": "keep"}],
                },
                schema=source.schema,
            ),
        ]
    ).with_columns(pl.concat_list("record").alias("nested"))
    before = source.clone()
    frame = source.lazy() if lazy else source
    engine = PolarsEngine()
    original_schema = engine.schema(frame)
    page = engine.page(frame, 0, 6)
    expected_record = {"when": raw, "ticks": 41, "text": "keep"}
    expected = [[raw], [raw, None], expected_record, [expected_record]]
    assert [cell["raw"] for cell in page["rows"][0]["values"]] == expected
    assert page["rows"][0]["values"] == page["rows"][2]["values"]
    assert page["rows"][0]["values"] != page["rows"][1]["values"]
    assert [cell["raw"] for cell in page["rows"][4]["values"]] == [[], None, None, [None]]
    assert [cell["raw"] for cell in page["rows"][5]["values"]] == [
        None,
        [None, None],
        {"when": None, "ticks": 41, "text": "keep"},
        [{"when": None, "ticks": 41, "text": "keep"}],
    ]
    for position, summary in enumerate(engine.summaries(frame)):
        label = json.dumps(expected[position], ensure_ascii=False, separators=(",", ":"))
        assert {item["value"]: item["count"] for item in summary["topValues"]}[label] == 2
        assert summary["distinctCount"] == source[:, position].drop_nulls().n_unique()
        assert summary["nullCount"] == source[:, position].null_count()
        assert all(item["selectionValue"] is None for item in summary["topValues"])
    choices, has_more = engine.column_values(frame.head(3), "record")
    record_label = json.dumps(expected_record, ensure_ascii=False, separators=(",", ":"))
    assert {item["value"]: item["count"] for item in choices}[record_label] == 2
    assert not has_more and all(item["selectionValue"] is None for item in choices)
    with pytest.raises(pl.exceptions.InvalidOperationError, match="cannot cast List"):
        engine.column_values(frame, "items")
    for empty in (source.head(0), source.slice(4, 1)):
        selected = empty.lazy() if lazy else empty
        assert len(engine.page(selected, 0, 2)["rows"]) == empty.height
        assert all(summary["totalCount"] == empty.height for summary in engine.summaries(selected))
    assert engine.schema(frame) == original_schema
    assert source.equals(before) and source.schema == before.schema
    json.dumps(page, allow_nan=False)
    engine.close()


def _literal_nested_temporal_source() -> tuple[pl.DataFrame, list[list[Any]]]:
    dtype = pl.Struct(
        {
            "^a.*$": pl.Datetime("ns"),
            "amount": pl.Datetime("ns"),
            "*": pl.Struct({"^a.*$": pl.Datetime("ns"), "amount": pl.Datetime("ns"), "__ow_field_0": pl.String}),
            "__ow_field_0": pl.String,
            "__ow_field_1": pl.Int64,
        }
    )
    record = {
        "^a.*$": -1,
        "amount": 1_000_000_001,
        "*": {"^a.*$": 1, "amount": -1, "__ow_field_0": "inner"},
        "__ow_field_0": "outer",
        "__ow_field_1": 23,
    }
    child_nulls = {**record, "^a.*$": None, "*": None}
    source = pl.DataFrame(
        {
            "items": pl.Series([[record], [record], [], None]).cast(pl.List(dtype)),
            "record": pl.Series([record, record, child_nulls, None]).cast(dtype),
            "fixed": pl.Series([[record, None], [record, None], [child_nulls, None], None]).cast(pl.Array(dtype, 2)),
        }
    )
    expected = {
        "^a.*$": "1969-12-31T23:59:59.999999999",
        "amount": "1970-01-01T00:00:01.000000001",
        "*": {
            "^a.*$": "1970-01-01T00:00:00.000000001",
            "amount": "1969-12-31T23:59:59.999999999",
            "__ow_field_0": "inner",
        },
        "__ow_field_0": "outer",
        "__ow_field_1": 23,
    }
    expected_nulls = {**expected, "^a.*$": None, "*": None}
    return source, [
        [[expected], expected, [expected, None]],
        [[expected], expected, [expected, None]],
        [[], expected_nulls, [expected_nulls, None]],
        [None, None, None],
    ]


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_nested_temporal_fields_are_literal(lazy: bool) -> None:
    source, expected = _literal_nested_temporal_source()
    before = source.clone()
    frame = source.lazy() if lazy else source
    engine = PolarsEngine()
    try:
        schema = engine.schema(frame)
        page = engine.page(frame, 0, 4)
        assert [[cell["raw"] for cell in row["values"]] for row in page["rows"]] == expected
        assert [row["rowNumber"] for row in page["rows"]] == list(range(4))
        for row, values in zip(page["rows"], expected, strict=True):
            assert [cell["display"] for cell in row["values"]] == [
                "" if value is None else json.dumps(value, ensure_ascii=False, separators=(",", ":"))
                for value in values
            ]
        for position, summary in enumerate(engine.summaries(frame)):
            label = json.dumps(expected[0][position], ensure_ascii=False, separators=(",", ":"))
            assert {value["value"]: value["count"] for value in summary["topValues"]}[label] == 2
            assert summary["nullCount"] == source[:, position].null_count()
            assert summary["distinctCount"] == source[:, position].drop_nulls().n_unique()
            assert summary["totalCount"] == 4
            assert all(value["selectionValue"] is None for value in summary["topValues"])
        assert engine.page(frame.head(0), 0, 4)["rows"] == []
        assert all(summary["totalCount"] == 0 for summary in engine.summaries(frame.head(0)))
        choices_source = pl.DataFrame(
            {
                "choice": pl.Series(
                    [
                        {"^a.*$": -1, "amount": 1_000_000_001, "text": "keep"},
                        {"^a.*$": -1, "amount": 1_000_000_001, "text": "keep"},
                        {"^a.*$": None, "amount": None, "text": "keep"},
                    ]
                ).cast(pl.Struct({"^a.*$": pl.Datetime("ns"), "amount": pl.Datetime("ns"), "text": pl.String}))
            }
        )
        choices_frame = choices_source.lazy() if lazy else choices_source
        choices, has_more = engine.column_values(choices_frame.head(2), "choice")
        label = json.dumps(
            {"^a.*$": "1969-12-31T23:59:59.999999999", "amount": "1970-01-01T00:00:01.000000001", "text": "keep"},
            separators=(",", ":"),
        )
        assert {value["value"]: value["count"] for value in choices} == {label: 2}
        assert not has_more and all(value["selectionValue"] is None for value in choices)
        with pytest.raises(pl.exceptions.InvalidOperationError, match="conversion from `struct"):
            engine.column_values(choices_frame, "choice")
        assert engine.schema(frame) == schema
    finally:
        engine.close()
    assert source.equals(before) and source.schema == before.schema


def test_polars_nontemporal_containers_do_not_enter_temporal_value_traversal(monkeypatch: pytest.MonkeyPatch) -> None:
    source = pl.DataFrame(
        {
            "items": pl.Series([[1, None], [], None], dtype=pl.List(pl.Int64)),
            "record": pl.Series([{"*": 1, "^a.*$": "keep"}, None, {"*": None, "^a.*$": "last"}]),
            "unselected": _literal_nested_temporal_source()[0].get_column("record").head(3),
        }
    )

    def refuse(*_args: Any, **_kwargs: Any) -> Any:
        raise AssertionError("Non-temporal values must not enter temporal formatting or decoding.")

    monkeypatch.setattr(polars_engine, "_polars_query_text", refuse)
    monkeypatch.setattr(polars_engine, "_polars_nested_temporal_value", refuse)
    ordinary = source.select("items", "record")
    assert polars_engine._polars_prepare_temporal_cells(ordinary, ordinary.schema) is ordinary
    engine = PolarsEngine()
    projection = [(0, "selected:items"), (1, "selected:record")]
    for frame in (source, source.lazy()):
        page = engine.page(frame, 0, 2, column_projection=projection)
        assert page["columnIds"] == [identifier for _, identifier in projection] and len(page["rows"]) == 2
        assert [cell["raw"] for cell in page["rows"][0]["values"]] == [[1, None], {"*": 1, "^a.*$": "keep"}]
        assert [summary["distinctCount"] for summary in engine.summaries(frame, projection)] == [2, 2]
    engine.close()


def test_polars_sliced_nested_page_formats_only_returned_children_in_one_batch(monkeypatch: pytest.MonkeyPatch) -> None:
    chunk = pl.DataFrame({"items": [[-1, None]]}).cast({"items": pl.List(pl.Datetime("ns"))})
    source = pl.concat([chunk] * 20, rechunk=False)
    before = source.clone()
    calls: list[int] = []
    native_text = polars_engine._polars_query_text

    def observe(values: pl.Series) -> pl.Series:
        calls.append(len(values))
        return values

    def tracked_text(expression: Any, dtype: Any) -> Any:
        return native_text(expression, dtype).map_batches(observe, return_dtype=pl.String, is_elementwise=True)

    monkeypatch.setattr(polars_engine, "_polars_query_text", tracked_text)
    engine = PolarsEngine()
    page = engine.page(source.lazy(), 3, 8)
    assert calls == [16]
    assert [row["rowNumber"] for row in page["rows"]] == list(range(3, 11))
    assert all(row["values"][0]["raw"] == ["1969-12-31T23:59:59.999999999", None] for row in page["rows"])
    assert source.equals(before) and source.schema == before.schema
    engine.close()


def test_polars_nested_temporal_file_session_keeps_native_generated_export_and_source(tmp_path: Path) -> None:
    source, expected_rows = _literal_nested_temporal_source()
    before = source.clone()
    path = tmp_path / "nested.parquet"
    source.write_parquet(path)
    contents, stat = path.read_bytes(), path.stat()
    manager = SessionManager()
    try:
        opened = manager.open_session({"kind": "file", "path": str(path)}, backend="polars", mode="editing")
        assert [[cell["raw"] for cell in row["values"]] for row in opened["page"]["rows"]] == expected_rows
        sid = opened["metadata"]["sessionId"]
        reference = {key: opened["metadata"]["schema"][0][key] for key in ("id", "name")}
        preview = manager.preview_step(
            sid, 0, {"id": "copy", "kind": "cloneColumn", "params": {"column": reference, "newName": "copy"}}, 0, 4
        )
        applied = manager.apply_draft(sid, preview["revision"], 0, 4)
        assert [row["id"] for row in applied["page"]["rows"]] == [row["id"] for row in opened["page"]["rows"]]
        assert [[cell["raw"] for cell in row["values"][:-1]] for row in applied["page"]["rows"]] == expected_rows
        assert [row["values"][-1] for row in applied["page"]["rows"]] == [
            row["values"][0] for row in opened["page"]["rows"]
        ]
        namespace: dict[str, Any] = {}
        exec(applied["code"], namespace)
        expected = source.with_columns(pl.col("items").alias("copy"))
        generated = namespace["clean_data"](source.lazy()).collect()
        assert generated.equals(expected) and generated.schema == expected.schema
        output = tmp_path / "export.parquet"
        output.touch()
        device, inode = _regular_file_identity(output)
        manager.export_data(
            sid, applied["revision"], str(output), {"format": "parquet"}, {"device": str(device), "inode": str(inode)}
        )
        exported = pl.read_parquet(output)
        assert exported.equals(expected) and exported.schema == expected.schema
    finally:
        manager.close_all()
    assert not manager.sessions and source.equals(before) and source.schema == before.schema
    after = path.stat()
    assert path.read_bytes() == contents
    assert (after.st_dev, after.st_ino, after.st_size, after.st_mtime_ns) == (
        stat.st_dev,
        stat.st_ino,
        stat.st_size,
        stat.st_mtime_ns,
    )


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize(
    "dtype,tick,raw,text,portable",
    [
        (pl.Duration("ns"), 10**17 + 1, "100000000.000000001", "1157d 9h 46m 40s 1ns", False),
        (pl.Datetime("ns"), 10**17 + 1, "1973-03-03T09:46:40.000000001", None, False),
        (pl.Duration("us"), 100000001, 100.000001, "1m 40s 1µs", True),
        (pl.Duration("us"), 1, "0.000001", "1µs", True),
        (pl.Duration("ms"), -12345, -12.345, "-12s -345ms", True),
        (pl.Datetime("ms"), -12345, "1969-12-31T23:59:47.655000", None, True),
        (pl.Datetime("us", "Europe/Berlin"), 100000001, "1970-01-01T01:01:40.000001+01:00", None, True),
        (pl.Datetime("ns", "Europe/Amsterdam"), -5364662400000000000, "1800-01-01T00:17:30+00:17:30", None, False),
        (pl.Duration("ns"), -(2**63), "-9223372036.854775808", "-106751d -23h -47m -16s -854775808ns", False),
        (pl.Datetime("ns"), -(2**63), "1677-09-21T00:12:43.145224192", None, False),
    ],
)
def test_polars_temporal_queries_preserve_native_values_before_row_boxing(
    monkeypatch: pytest.MonkeyPatch, lazy: bool, dtype: Any, tick: int, raw: Any, text: str | None, portable: bool
) -> None:
    import __main__

    ticks = [tick, tick + 1, tick, 0, None]
    source = pl.DataFrame({"value": pl.Series(ticks, dtype=pl.Int64).cast(dtype), "row": range(5)})
    before = source.clone()
    frame = source.lazy() if lazy else source
    monkeypatch.setattr(__main__, "exact_polars_temporal_source", frame, raising=False)
    engine = PolarsEngine()
    manager = SessionManager()
    view = {"filters": [], "sort": []}
    try:
        opened = manager.open_session(
            {"kind": "notebookVariable", "label": "temporal values", "variableName": "exact_polars_temporal_source"},
            backend="polars",
            page_size=5,
        )
        metadata = opened["metadata"]
        sid, revision = metadata["sessionId"], metadata["revision"]
        schema = metadata["schema"]
        kind = "duration" if isinstance(dtype, pl.Duration) else "datetime"
        cell = {"kind": kind, "raw": raw, "display": text or raw, "isNull": False, "isNaN": False}
        cells = [row["values"][0] for row in opened["page"]["rows"]]
        assert cells[0] == cells[2] == cell
        assert cells[0]["raw"] != cells[1]["raw"] and cells[0]["display"] != cells[1]["display"]
        assert cells[-1]["isNull"] and cells[-1]["raw"] is None
        choices = manager.get_column_values(sid, revision, "value", view, limit=1)
        choice = choices["values"][0]
        assert choice["value"] == cell["display"] and choice["count"] == 2
        assert len(choices["values"]) == 1
        searched = manager.get_column_values(sid, revision, "value", view, search=cell["display"], limit=1)
        assert searched["values"] == [choice]
        if kind == "datetime":
            alias = str(cell["display"]).replace("T", " ")
            assert manager.get_column_values(sid, revision, "value", view, search=alias)["values"] == [choice]
        summary = manager.get_summary(sid, revision, view, [schema[0]["id"]])["summaries"][0]
        assert (summary["nullCount"], summary["nanCount"], summary["distinctCount"]) == (1, 0, 3)
        assert summary["topValues"][0] == choice
        if kind == "datetime":
            present = [(value, cells[index]["display"]) for index, value in enumerate(ticks) if value is not None]
            assert summary["visualization"]["min"] == min(present)[1]
            assert summary["visualization"]["max"] == max(present)[1]
        token = {"kind": "typedSelection", "version": 1, "columnType": kind, "cell": cell}
        if portable:
            assert choice["selectionValue"] == token
        else:
            assert choice["selectionValue"] is None
        column_filter = {
            "column": "value",
            "type": kind,
            "predicates": [],
            "valueFilter": {
                "kind": "values",
                "selectedValues": [summary["topValues"][0]["selectionValue"] if portable else token],
                "includeNulls": False,
                "includeNaN": False,
            },
        }
        model = {"filters": [column_filter], "sort": []}
        lineage = source_lineage(engine.schema(frame))
        step = bind_step(
            validate_step(
                {
                    "id": "exact-temporal-filter",
                    "kind": "filterRows",
                    "params": {"filterModel": {"filters": [{**column_filter, "column": lineage[0]}], "sort": []}},
                }
            ),
            engine.schema(frame),
            lineage,
        )
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([step]), namespace)
        if portable:
            selected = manager.get_page(sid, revision, 0, 5, model)["page"]
            assert [row["values"][1]["raw"] for row in selected["rows"]] == [0, 2]
            for actual in [engine.apply_transform(frame, step), namespace["clean_data"](frame)]:
                result = actual.collect() if lazy else actual
                assert result.equals(source[[0, 2]])
        else:
            with pytest.raises(EngineError):
                manager.get_page(sid, revision, 0, 5, model)
            with pytest.raises((EngineError, ValueError)):
                engine.apply_transform(frame, step)
            with pytest.raises(ValueError):
                namespace["clean_data"](frame)
        restored = manager.get_page(sid, revision, 0, 5, view)
        assert [row["values"][1]["raw"] for row in restored["page"]["rows"]] == list(range(5))
        assert all(restored["metadata"][key] == metadata[key] for key in ("sessionId", "revision", "source"))
        for empty in [source.head(0), source.tail(1)]:
            empty_frame = empty.lazy() if lazy else empty
            assert engine.column_values(empty_frame, "value") == ([], False)
            assert engine.summaries(empty_frame, [(0, "value")])[0]["nullCount"] == empty.height
        assert source.equals(before) and source["value"].cast(pl.Int64).to_list() == ticks
        json.dumps(opened, allow_nan=False)
    finally:
        manager.close_all()
        engine.close()


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize("unit,scale", [("ms", 1_000), ("us", 1_000_000), ("ns", 1_000_000_000)])
def test_polars_datetime_choices_label_like_cells_and_search_either_separator(
    lazy: bool, unit: Any, scale: int
) -> None:
    whole, fraction = "2020-01-01T00:00:00", "2020-01-01T00:00:00.123000"
    source = pl.DataFrame(
        {
            "value": pl.Series([1577836800 * scale, 1577836800 * scale + 123 * (scale // 1000), None]).cast(
                pl.Datetime(unit)
            )
        }
    )
    frame = source.lazy() if lazy else source
    engine = PolarsEngine()
    try:
        choices, truncated = engine.column_values(frame, "value")
        assert not truncated and [item["value"] for item in choices] == [whole, fraction]
        assert [item["selectionValue"]["cell"]["raw"] for item in choices] == [whole, fraction]
        for needle, expected in [
            (".000", []),
            ("00:00:00", choices),
            (".123000", [choices[1]]),
            (".123000000", []),
            (fraction, [choices[1]]),
            (fraction.replace("T", " "), [choices[1]]),
            (fraction.replace("T", "t"), [choices[1]]),
        ]:
            assert engine.column_values(frame, "value", search=needle)[0] == expected
    finally:
        engine.close()


@pytest.mark.parametrize("label", ["integer", "decimal", "bool", "array", "struct", "datetime", "plain"])
@pytest.mark.parametrize("lazy", [False, True])
def test_polars_enum_labels_do_not_change_profiles_or_typed_filters(label: str, lazy: bool) -> None:
    engine = PolarsEngine()
    source = pl.DataFrame({"value": pl.Series([label, "other", None], dtype=pl.Enum([label, "other"]))})
    before = source.clone()
    frame = source.lazy() if lazy else source
    assert engine.schema(frame)[0]["type"] == "string"
    summary = engine.summaries(frame)[0]
    assert summary["type"] == "string"
    assert summary["nullCount"] == 1
    values, truncated = engine.column_values(frame, "value")
    selected = next(item["selectionValue"] for item in values if item["value"] == label)
    assert not truncated
    assert selected == typed_selection_value(label, "string")
    column_filter = {
        "column": "value",
        "type": "string",
        "predicates": [],
        "valueFilter": {"kind": "values", "selectedValues": [selected], "includeNulls": False, "includeNaN": False},
    }
    filtered = engine.apply_filter_model(frame, {"filters": [column_filter], "sort": []})
    assert isinstance(filtered, pl.LazyFrame) == lazy
    assert (filtered.collect() if lazy else filtered).rows() == [(label,)]
    schema = engine.schema(frame)
    lineage = source_lineage(schema)
    operation = bind_step(
        validate_step(
            {
                "id": "enum-filter",
                "kind": "filterRows",
                "params": {
                    "filterModel": {
                        "filters": [{**column_filter, "column": lineage[0]}],
                        "sort": [],
                    }
                },
            }
        ),
        schema,
        lineage,
    )
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([operation]), namespace)
    generated = namespace["clean_data"](frame)
    assert isinstance(generated, pl.LazyFrame) == lazy
    assert (generated.collect() if lazy else generated).rows() == [(label,)]
    assert source.equals(before)


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_nested_array_does_not_replace_outer_struct_family(lazy: bool) -> None:
    engine = PolarsEngine()
    source = pl.DataFrame(
        {"value": pl.Series([{"payload": [1, 2]}, None], dtype=pl.Struct({"payload": pl.Array(pl.Int64, 2)}))}
    )
    before = source.clone()
    frame = source.lazy() if lazy else source
    assert engine.schema(frame)[0]["type"] == "struct"
    summary = engine.summaries(frame)[0]
    assert summary["type"] == "struct"
    assert summary["nullCount"] == 1
    with pytest.raises(EngineError, match="sorting is unavailable for struct columns"):
        engine.apply_filter_model(frame, {"filters": [], "sort": [{"column": "value", "direction": "asc"}]})
    assert source.equals(before)


def reserve_export_target(path: Path) -> dict[str, str]:
    path.touch(exist_ok=False)
    device, inode = _regular_file_identity(path)
    return {"device": str(device), "inode": str(inode)}


def export_options(format_name: str) -> dict[str, object]:
    return (
        {"format": "csv", "delimiter": ",", "quoteChar": '"', "encoding": "utf-8", "header": True}
        if format_name == "csv"
        else {"format": "parquet"}
    )


def _write_polars_file(path: Path, extension: str, values: list[int]) -> None:
    frame = pl.DataFrame({"value": values, "label": [f"row-{index}" for index in range(len(values))]})
    if extension == "csv":
        frame.write_csv(path)
    elif extension == "tsv":
        frame.write_csv(path, separator="\t")
    elif extension == "parquet":
        frame.write_parquet(path)
    else:
        assert extension == "jsonl"
        frame.write_ndjson(path)


@pytest.mark.parametrize(
    ("source", "pyarrow_available", "expected_imports"),
    (
        ({"kind": "file", "path": "sample.XLSX"}, True, ["polars", "pyarrow"]),
        ({"kind": "file", "path": "legacy.xls"}, True, ["polars", "pyarrow"]),
        ({"kind": "file", "path": "sample.xlsx"}, False, ["polars"]),
        ({"kind": "file", "path": "sample.csv"}, True, ["polars"]),
        ({"kind": "notebookVariable", "variableName": "frame"}, True, ["polars"]),
    ),
)
def test_polars_preparation_preloads_only_an_installed_excel_pyarrow_bridge(
    source: dict[str, Any],
    pyarrow_available: bool,
    expected_imports: list[str],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    imported: list[str] = []

    def find_optional_module(name: str) -> object | None:
        assert name == "pyarrow"
        return object() if pyarrow_available else None

    monkeypatch.setattr(engine_base, "import_module", imported.append)
    monkeypatch.setattr(polars_engine, "find_spec", find_optional_module)
    monkeypatch.setattr(polars_engine, "import_module", imported.append)

    PolarsEngine().prepare(source)

    assert imported == expected_imports


def test_polars_file_session_pages_filters_and_summarizes_without_pandas(monkeypatch):
    import builtins

    original_import = builtins.__import__

    def reject_pandas_import(name, *args, **kwargs):
        if name == "pandas" or name.startswith("pandas."):
            raise AssertionError("Polars sessions must not import pandas")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", reject_pandas_import)

    def fail_to_pandas(*_args, **_kwargs):
        raise AssertionError("Polars sessions must not convert to pandas")

    monkeypatch.setattr(pl.DataFrame, "to_pandas", fail_to_pandas, raising=False)

    manager = SessionManager()
    opened = manager.open_session(
        {"kind": "file", "label": "sample.csv", "path": str(ROOT / "fixtures" / "sample.csv")},
        backend="polars",
        page_size=2,
    )

    assert opened["metadata"]["backend"] == "polars"
    assert opened["metadata"]["shape"] == {"rows": 4, "columns": 4}
    assert "stats" not in opened["metadata"]
    assert opened["page"]["rows"][0]["values"][0]["display"] == "Milan"

    session = manager.sessions[opened["metadata"]["sessionId"]]
    assert isinstance(session.original, pl.LazyFrame)
    stats = manager.get_dataset_stats(opened["metadata"]["sessionId"], 0, {"filters": [], "sort": []})
    assert stats["stats"]["missingValuesByColumn"][0] == {"column": "city", "count": 0}

    filter_model = {
        "filters": [
            {
                "column": "city",
                "type": "string",
                "valueFilter": None,
                "predicates": [{"kind": "predicate", "operator": "contains", "value": "i"}],
            }
        ],
        "sort": [{"column": "sales", "direction": "desc", "nulls": "last"}],
    }
    page = manager.get_page(opened["metadata"]["sessionId"], 0, 0, 10, filter_model)

    assert page["metadata"]["filteredShape"]["rows"] == 3
    assert [row["values"][0]["display"] for row in page["page"]["rows"]] == ["Berlin", "Milan", "Paris"]

    sales_id = next(column["id"] for column in opened["metadata"]["schema"] if column["name"] == "sales")
    summary = manager.get_summary(opened["metadata"]["sessionId"], 0, filter_model, [sales_id])
    assert summary["summaries"][0]["numeric"]["max"] == 12.0
    assert summary["summaries"][0]["visualization"]["kind"] == "numeric"
    assert summary["summaries"][0]["visualization"]["bins"]


@pytest.mark.parametrize("extension", ["csv", "tsv", "parquet", "jsonl"])
def test_polars_file_scans_treat_glob_metacharacters_as_literal_path_characters(
    extension: str,
    tmp_path: Path,
) -> None:
    path = tmp_path / f"[published] source.{extension}"
    _write_polars_file(path, extension, [17, 18])
    _write_polars_file(tmp_path / f"p source.{extension}", extension, [99])

    options = {"delimiter": "\t"} if extension == "tsv" else None
    if extension == "jsonl" and os.name == "nt":
        with pytest.raises(EngineError, match="Windows.*glob"):
            PolarsEngine().read_file(str(path), options)
        return
    frame = PolarsEngine().read_file(str(path), options)

    assert isinstance(frame, pl.LazyFrame)
    assert frame.collect().get_column("value").to_list() == [17, 18]


@pytest.mark.parametrize(
    ("suffix", "delimiter", "record_ending", "line_ending"),
    [
        ("csv", ";", "\r", "cr"),
        ("tsv", "\t", "\r", "cr"),
        ("csv", ";", "\n", None),
        ("csv", ";", "\r\n", None),
        ("csv", ";", "\r\n", "lf"),
    ],
    ids=["cr-csv", "cr-tsv", "lf-omitted", "crlf-omitted", "crlf-explicit"],
)
@pytest.mark.parametrize("has_header", [True, False], ids=["header", "headerless"])
def test_polars_delimited_line_endings_preserve_native_records(
    tmp_path: Path, suffix: str, delimiter: str, record_ending: str, line_ending: str | None, has_header: bool
) -> None:
    path = tmp_path / f"records.{suffix}"
    records = [f'"Milan\n""centre"""{delimiter}1', f'"Berlin\r\nwest"{delimiter}2']
    if has_header:
        records.insert(0, f"city{delimiter}value")
    contents = (record_ending.join(records) + record_ending).encode("utf-8")
    path.write_bytes(contents)
    options: dict[str, Any] = {"quoteChar": '"', "hasHeader": has_header}
    if suffix == "csv":
        options["delimiter"] = delimiter
    if line_ending is not None:
        options["lineEnding"] = line_ending
    manager = SessionManager()
    try:
        opened = manager.open_session(
            {"kind": "file", "label": path.name, "path": str(path), "importOptions": options},
            backend="polars",
            page_size=1,
        )
        metadata = opened["metadata"]
        session_id = metadata["sessionId"]
        assert metadata["shape"] == {"rows": 2, "columns": 2}
        assert metadata["source"]["importOptions"] == options
        assert [column["name"] for column in metadata["schema"]] == (
            ["city", "value"] if has_header else ["column_1", "column_2"]
        )
        assert isinstance(manager.sessions[session_id].original, pl.LazyFrame)
        assert [cell["raw"] for cell in opened["page"]["rows"][0]["values"]] == ['Milan\n"centre"', 1]
        later = manager.get_page(session_id, 0, 1, 1, {"filters": [], "sort": []})
        assert [cell["raw"] for cell in later["page"]["rows"][0]["values"]] == ["Berlin\r\nwest", 2]
    finally:
        manager.close_all()
    assert manager.sessions == {}
    assert path.read_bytes() == contents


@pytest.mark.parametrize(
    ("name", "verbatim"),
    [
        *[
            pytest.param(name, False, marks=pytest.mark.skipif(os.name == "nt", reason="Unix literal-path read."))
            for name in [
                "[selected].jsonl",
                "question?.ndjson",
                "star*.ndjson",
                "space %20 {x}].ndjson",
                "[nested]/source.ndjson",
                "back\\slash.ndjson",
            ]
        ],
        *[
            pytest.param(name, True, marks=pytest.mark.skipif(os.name != "nt", reason="Native Windows path read."))
            for name in ["plain.jsonl", "space %20 {x}].ndjson"]
        ],
    ],
)
def test_polars_ndjson_session_reads_only_the_selected_file_and_invalidates_replacement(
    name: str, verbatim: bool, tmp_path: Path
) -> None:
    path = tmp_path / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b'{"value":17}\n{"value":18}\n')
    before = path.read_bytes()
    escaped = Path(glob.escape(str(path)))
    if escaped != path:
        escaped.mkdir(parents=True)
        (escaped / "other.ndjson").write_bytes(b'{"value":99}\n')
    encoded = tmp_path / path.as_uri().rsplit("/", 1)[1]
    if encoded != path:
        encoded.write_bytes(b'{"value":99}\n')
    selected_path = "\\\\?\\" + str(path.absolute()) if verbatim else str(path)
    manager = SessionManager()
    opened = manager.open_session({"kind": "file", "label": name, "path": selected_path}, backend="polars", page_size=1)
    session_id = opened["metadata"]["sessionId"]
    try:
        assert opened["page"]["rows"][0]["values"][0]["display"] == "17"
        assert opened["metadata"]["shape"] == {"rows": 2, "columns": 1}
        page = manager.get_page(session_id, 0, 1, 1, {"filters": [], "sort": []})
        assert page["page"]["rows"][0]["values"][0]["display"] == "18"
        frame = manager.sessions[session_id].original
        engine = PolarsEngine()
        schema = engine.schema(frame)
        lineage = source_lineage(schema)
        operation = bind_step(
            validate_step(
                {
                    "id": "sort",
                    "kind": "sortRows",
                    "params": {"rules": [{"column": lineage[0], "direction": "desc", "nulls": "last"}]},
                }
            ),
            schema,
            lineage,
        )
        engine.validate_transform_preflight(frame, operation, engine.shape(frame))
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan([operation]), namespace)
        for result in (engine.apply_transform(frame, operation), namespace["clean_data"](frame)):
            assert isinstance(result, pl.LazyFrame)
            assert result.collect().get_column("value").to_list() == [18, 17]
        assert path.read_bytes() == before
        path.rename(tmp_path / "original.ndjson")
        path.write_bytes(b'{"value":99}\n')
        with pytest.raises(SourceChangedError):
            manager.get_page(session_id, 0, 0, 1, {"filters": [], "sort": []})
    finally:
        manager.close_session(session_id, 0)
    assert not manager.sessions


@pytest.mark.skipif(os.name == "nt", reason="The native descriptor bridge is Unix-only.")
def test_polars_ndjson_native_plan_outlives_the_closed_builtin_stream(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "lifetime.ndjson"
    path.write_bytes(b'{"value":17}\n{"value":18}\n')
    native_scan = pl.scan_ndjson
    streams: list[io.BufferedReader] = []

    def scan(source: Any) -> pl.LazyFrame:
        assert type(source) is io.BufferedReader
        assert not source.closed
        streams.append(source)
        return native_scan(source)

    monkeypatch.setattr(pl, "scan_ndjson", scan)
    frame = PolarsEngine().read_file(str(path))
    assert streams[0].closed
    clone = frame.clone()
    del frame
    assert clone.select("value").limit(1).collect().to_dicts() == [{"value": 17}]
    assert clone.collect().get_column("value").to_list() == [17, 18]


@pytest.mark.parametrize(
    ("name", "accepted"),
    [
        (r"C:\plain.ndjson", True),
        (r"C:\space %20 {x}].jsonl", True),
        (r"C:\[selected].jsonl", False),
        (r"C:\question?.jsonl", False),
        (r"C:\star*.jsonl", False),
        (r"\\?\C:\plain.jsonl", True),
        (r"\\?\c:\space %20 {x}].ndjson", True),
        (r"\\?\C:\[selected].jsonl", False),
        (r"\\?\C:\[parent]\plain.jsonl", False),
        (r"\\?\C:\question?.jsonl", False),
        (r"\\?\C:\star*.jsonl", False),
        (r"\\?\UNC\server\share\plain.jsonl", False),
        (r"\\server\share\plain.jsonl", True),
        (r"\\?\Volume{fixture}\plain.jsonl", False),
        (r"\\?\C:plain.jsonl", False),
    ],
)
def test_polars_ndjson_windows_branch_uses_direct_paths_or_refuses_glob_syntax(
    name: str, accepted: bool, monkeypatch: pytest.MonkeyPatch
) -> None:
    # This checks dispatch on every platform; native Windows qualification is separate.
    class WindowsPath(PureWindowsPath):
        def expanduser(self) -> WindowsPath:
            return self

        def absolute(self) -> WindowsPath:
            return self

    calls: list[str] = []
    monkeypatch.setattr(polars_engine, "Path", WindowsPath)
    monkeypatch.setattr(polars_engine, "os", SimpleNamespace(name="nt"))
    monkeypatch.setattr(pl, "scan_ndjson", lambda source: calls.append(source))
    if not accepted:
        # pathlib versions classify this malformed prefix at different refusal boundaries.
        message = (
            None
            if name == r"\\?\C:plain.jsonl"
            else r"openWrangler\.defaultBackend.*pandas.*Open Wrangler: Open File Path"
        )
        with pytest.raises(EngineError, match=message):
            PolarsEngine().read_file(name)
        assert not calls
    else:
        PolarsEngine().read_file(name)
        assert calls == [name]


@pytest.mark.skipif(os.name == "nt", reason="Resource limits and native file duplication are Unix-only.")
@pytest.mark.parametrize("exhaust", [False, True])
@pytest.mark.parametrize("constructor_failure", [False, True])
def test_polars_ndjson_native_handle_failure_refuses_buffering_in_an_isolated_process(
    exhaust: bool, constructor_failure: bool, tmp_path: Path
) -> None:
    path = tmp_path / "source.ndjson"
    path.write_bytes(b'{"value":17}\n' * 100_000)
    script = dedent("""
        import gc, os, resource, sys, tracemalloc
        from pathlib import Path
        import polars as pl
        from openwrangler_runtime.engines.polars_engine import PolarsEngine
        from openwrangler_runtime.engines.base import EngineError

        path, exhaust, constructor_failure = sys.argv[1], sys.argv[2] == 'True', sys.argv[3] == 'True'
        engine = PolarsEngine()
        engine.read_file(path).limit(1).collect()
        gc.collect()
        native_scan = pl.scan_ndjson
        if constructor_failure:
            def fail_after_source(*args, **kwargs):
                native_scan(*args, **kwargs)
                raise ValueError('owned constructor failure')
            pl.scan_ndjson = fail_after_source
        limit = resource.getrlimit(resource.RLIMIT_NOFILE)
        fillers = []
        frame = None
        outcome = 'returned'
        try:
            if exhaust:
                soft = 64 if limit[0] == resource.RLIM_INFINITY else min(64, limit[0])
                resource.setrlimit(resource.RLIMIT_NOFILE, (soft, limit[1]))
                while True:
                    try:
                        fillers.append(os.open(os.devnull, os.O_RDONLY))
                    except OSError:
                        break
                os.close(fillers.pop())  # Builtin open succeeds, but native duplication cannot.
            tracemalloc.start()
            try:
                frame = engine.read_file(path)
            except EngineError as error:
                assert 'native file handle' in str(error), str(error)
                outcome = 'refused'
            except ValueError as error:
                assert str(error) == 'owned constructor failure'
                outcome = 'constructor failure'
            finally:
                peak = tracemalloc.get_traced_memory()[1]
                tracemalloc.stop()
        finally:
            for descriptor in fillers:
                os.close(descriptor)
            resource.setrlimit(resource.RLIMIT_NOFILE, limit)
        expected = 'refused' if exhaust else 'constructor failure' if constructor_failure else 'returned'
        assert outcome == expected, (outcome, expected)
        assert peak < Path(path).stat().st_size // 2, peak
        if frame is not None:
            assert frame.limit(1).collect().to_dicts() == [{'value': 17}]
        del frame
        gc.collect()
        if sys.platform.startswith('linux'):
            selected = Path(path).stat()
            for entry in Path('/proc/self/fd').iterdir():
                try:
                    metadata = os.fstat(int(entry.name))
                except OSError:
                    continue
                assert (metadata.st_dev, metadata.st_ino) != (selected.st_dev, selected.st_ino)
    """)
    result = subprocess.run(
        [sys.executable, "-c", script, str(path), str(exhaust), str(constructor_failure)],
        env={**os.environ, "PYTHONPATH": str(ROOT / "python")},
        capture_output=True,
        text=True,
        timeout=20,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert result.stderr == ""


def test_polars_literal_file_scan_disables_a_native_glob_option(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "[published] source.parquet"
    _write_polars_file(path, "parquet", [17, 18])
    native_scan_parquet = pl.scan_parquet
    calls: list[tuple[str, bool]] = []

    def scan_parquet(source: str, *, glob: bool = True, **options: Any) -> pl.LazyFrame:
        calls.append((source, glob))
        return native_scan_parquet(source, glob=glob, **options)

    monkeypatch.setattr(pl, "scan_parquet", scan_parquet)

    frame = PolarsEngine().read_file(str(path))

    assert isinstance(frame, pl.LazyFrame)
    assert frame.collect().height == 2
    assert calls == [(str(path), False)]


def test_polars_session_opens_pages_and_closes_a_literal_bracket_path(tmp_path: Path) -> None:
    path = tmp_path / "[Live] customer snapshot.csv"
    _write_polars_file(path, "csv", [17, 18])
    source_bytes = path.read_bytes()
    manager = SessionManager()

    opened = manager.open_session(
        {"kind": "file", "label": path.name, "path": str(path)},
        backend="polars",
        page_size=1,
    )
    session_id = opened["metadata"]["sessionId"]

    assert opened["metadata"]["shape"] == {"rows": 2, "columns": 2}
    assert isinstance(manager.sessions[session_id].original, pl.LazyFrame)
    second_page = manager.get_page(session_id, 0, 1, 1, {"filters": [], "sort": []})
    assert len(second_page["page"]["rows"]) == 1
    assert manager.close_session(session_id, 0) == {"kind": "sessionClosed", "sessionId": session_id}
    assert session_id not in manager.sessions
    assert path.read_bytes() == source_bytes


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_page_bounds_boxed_strings_without_changing_source_queries(
    monkeypatch: pytest.MonkeyPatch, lazy: bool
) -> None:
    exact = "🙂e\u0301\0" * 16_384
    oversized = "x" * 262_144 + "TAIL"
    source = pl.DataFrame(
        {
            "id": range(5),
            "*": [oversized, exact, None, "", oversized],
            "^a.*$": ["literal"] * 5,
            "apple": ["u" * 262_144] * 5,
        }
    )
    before = source.clone()
    boxed_lengths: list[int | None] = []
    boxed_columns: list[list[str]] = []
    native_iter_rows = pl.DataFrame.iter_rows

    def observe_rows(frame: pl.DataFrame, *args: Any, **kwargs: Any) -> Iterator[Any]:
        boxed_columns.append(frame.columns)
        for row in cast(Iterator[Any], native_iter_rows(frame, *args, **kwargs)):
            if isinstance(row, dict) and "*" in row:
                boxed_lengths.append(None if row["*"] is None else len(row["*"]))
            yield row

    monkeypatch.setattr(pl.DataFrame, "iter_rows", observe_rows)
    engine = PolarsEngine()
    try:
        frame = engine.ensure_row_ids(source.lazy() if lazy else source, "bounded-text")
        page = engine.page(frame, 1, 4, total_rows=5, column_projection=[(1, "star"), (2, "pattern")])

        # These are actual iter_rows values before typed-cell normalization.
        assert boxed_lengths == [65_536, None, 0, 65_537]
        assert len(boxed_columns) == 1 and boxed_columns[0][1:] == ["*", "^a.*$"]
        assert page["columnIds"] == ["star", "pattern"]
        assert [row["rowNumber"] for row in page["rows"]] == [1, 2, 3, 4]
        assert [row["id"].rsplit(":", 1)[-1] for row in page["rows"]] == ["1", "2", "3", "4"]
        cells = [row["values"][0] for row in page["rows"]]
        assert [cell["raw"] for cell in cells] == [exact, None, "", "x" * 65_537]
        assert [cell["display"] for cell in cells] == [exact, "", "", "x" * 65_537]
        assert [cell["kind"] for cell in cells] == ["string", "null", "string", "string"]
        assert [row["values"][1]["raw"] for row in page["rows"]] == ["literal"] * 4

        filtered = engine.apply_filter_model(
            frame,
            {
                "filters": [
                    {
                        "column": "*",
                        "type": "string",
                        "predicates": [{"kind": "predicate", "operator": "endsWith", "value": "TAIL"}],
                    }
                ],
                "sort": [{"column": "id", "direction": "asc", "nulls": "last"}],
            },
        )
        safe = engine.page(filtered, 0, 2, total_rows=2, column_projection=[(0, "id")])
        assert [row["values"][0]["raw"] for row in safe["rows"]] == [0, 4]
        assert len(boxed_columns) == 2 and boxed_columns[-1][1:] == ["id"]
        assert isinstance(frame, pl.LazyFrame) == isinstance(filtered, pl.LazyFrame) == lazy
        assert source.schema == before.schema and source.equals(before)
    finally:
        engine.close()


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_page_bounds_binary_boxing_when_native_slicing_is_available(
    monkeypatch: pytest.MonkeyPatch, lazy: bool
) -> None:
    exact = b"\x00\xff" * 24_576
    overflow = exact + b"\x00"
    oversized = b"\x00\xff" * 131_072
    source = pl.DataFrame(
        {
            "id": range(6),
            "*": pl.Series([oversized, exact, None, b"", overflow, oversized], dtype=pl.Binary),
            "^a.*$": [7] * 6,
            "apple": [oversized] * 6,
        }
    )
    before = source.clone()
    native_slice_available = callable(getattr(pl.col("*").bin, "slice", None))
    expected_large = overflow if native_slice_available else oversized
    boxed_values: list[bytes | None] = []
    boxed_columns: list[list[str]] = []
    expression_batches = 0
    native_iter_rows = pl.DataFrame.iter_rows
    native_with_columns = pl.DataFrame.with_columns

    def observe_rows(frame: pl.DataFrame, *args: Any, **kwargs: Any) -> Iterator[Any]:
        boxed_columns.append(frame.columns)
        for row in cast(Iterator[Any], native_iter_rows(frame, *args, **kwargs)):
            if isinstance(row, dict) and "*" in row:
                boxed_values.append(row["*"])
            yield row

    def observe_batch(frame: pl.DataFrame, *args: Any, **kwargs: Any) -> pl.DataFrame:
        nonlocal expression_batches
        expression_batches += 1
        return native_with_columns(frame, *args, **kwargs)

    monkeypatch.setattr(pl.DataFrame, "iter_rows", observe_rows)
    monkeypatch.setattr(pl.DataFrame, "with_columns", observe_batch)
    engine = PolarsEngine()
    try:
        frame = engine.ensure_row_ids(source.lazy() if lazy else source, "bounded-binary")
        page = engine.page(frame, 1, 5, total_rows=6, column_projection=[(1, "star"), (2, "pattern")])

        # Observe native bytes before base64; minimum Polars retains full boxing.
        assert [None if value is None else len(value) for value in boxed_values] == [
            49_152,
            None,
            0,
            49_153,
            len(expected_large),
        ]
        assert boxed_values == [exact, None, b"", overflow, expected_large]
        assert expression_batches == int(native_slice_available)
        assert len(boxed_columns) == 1 and boxed_columns[0][1:] == ["*", "^a.*$"]
        assert page["columnIds"] == ["star", "pattern"]
        assert [row["rowNumber"] for row in page["rows"]] == [1, 2, 3, 4, 5]
        assert [row["id"].rsplit(":", 1)[-1] for row in page["rows"]] == ["1", "2", "3", "4", "5"]
        cells = [row["values"][0] for row in page["rows"]]
        encoded = [
            b64encode(value).decode("ascii") if value is not None else None
            for value in [exact, None, b"", overflow, expected_large]
        ]
        assert [cell["raw"] for cell in cells] == encoded
        assert [cell["display"] for cell in cells] == [value or "" for value in encoded]
        assert [cell["kind"] for cell in cells] == ["binary", "null", "binary", "binary", "binary"]
        assert [row["values"][1]["raw"] for row in page["rows"]] == [7] * 5

        filtered = engine.apply_filter_model(
            frame,
            {
                "filters": [
                    {
                        "column": "*",
                        "type": "binary",
                        "predicates": [{"kind": "predicate", "operator": "isNotNull"}],
                    }
                ],
                "sort": [{"column": "id", "direction": "desc", "nulls": "last"}],
            },
        )
        safe = engine.page(filtered, 0, 5, total_rows=5, column_projection=[(0, "id")])
        assert [row["values"][0]["raw"] for row in safe["rows"]] == [5, 4, 3, 1, 0]
        assert len(boxed_columns) == 2 and boxed_columns[-1][1:] == ["id"]
        assert expression_batches == int(native_slice_available)
        assert isinstance(frame, pl.LazyFrame) == isinstance(filtered, pl.LazyFrame) == lazy
        assert source.schema == before.schema and source.equals(before)
    finally:
        engine.close()


def test_lazy_polars_page_projects_before_the_terminal_collect(monkeypatch: pytest.MonkeyPatch) -> None:
    engine = PolarsEngine()
    frame = engine.ensure_row_ids(
        pl.DataFrame({"omitted": [10, 20], "selected": [30, 40], "also_omitted": [50, 60]}).lazy(),
        "projection-order",
    )
    events: list[str] = []
    selected_columns: list[list[str]] = []
    native_select = pl.LazyFrame.select
    native_collect = pl.LazyFrame.collect

    def tracked_select(lazy_frame: pl.LazyFrame, *columns: Any, **kwargs: Any) -> pl.LazyFrame:
        events.append("select")
        values = columns[0] if len(columns) == 1 and isinstance(columns[0], list) else columns
        selected_columns.append([str(value) for value in values])
        return native_select(lazy_frame, *columns, **kwargs)

    def tracked_collect(lazy_frame: pl.LazyFrame, *args: Any, **kwargs: Any) -> pl.DataFrame:
        events.append("collect")
        return cast(pl.DataFrame, native_collect(lazy_frame, *args, **kwargs))

    monkeypatch.setattr(pl.LazyFrame, "select", tracked_select)
    monkeypatch.setattr(pl.LazyFrame, "collect", tracked_collect)

    page = engine.page(
        frame,
        0,
        2,
        total_rows=2,
        column_projection=[(1, "stable:selected")],
    )

    assert events == ["select", "collect"]
    assert selected_columns and "selected" in selected_columns[0]
    assert "omitted" not in selected_columns[0]
    assert "also_omitted" not in selected_columns[0]
    assert page["columnIds"] == ["stable:selected"]
    assert [row["values"][0]["display"] for row in page["rows"]] == ["30", "40"]


@pytest.mark.parametrize("extension", ["csv", "parquet", "jsonl"])
def test_real_polars_scan_selects_only_the_page_projection_before_collect(
    extension: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / f"projection.{extension}"
    source = pl.DataFrame({"omitted": [10, 20], "selected": [30, 40], "also_omitted": [50, 60]})
    if extension == "parquet":
        source = pl.DataFrame(
            {
                "omitted": range(257),
                "selected": pl.Series([10**17 + index + 1 for index in range(257)], dtype=pl.Int64).cast(
                    pl.Datetime("ns")
                ),
                "also_omitted": range(257),
            }
        )
    if extension == "csv":
        source.write_csv(path)
    elif extension == "parquet":
        source.write_parquet(path)
    else:
        source.write_ndjson(path)

    engine = PolarsEngine()
    frame = engine.ensure_row_ids(engine.read_file(str(path)), f"real-{extension}-projection")
    assert isinstance(frame, pl.LazyFrame)
    private_columns = set(frame.collect_schema().names()) - {"omitted", "selected", "also_omitted"}
    assert len(private_columns) == 1
    row_id = next(iter(private_columns))
    events: list[str] = []
    selected_columns: list[list[str]] = []
    native_select = pl.LazyFrame.select
    native_collect = pl.LazyFrame.collect

    def tracked_select(lazy_frame: pl.LazyFrame, *columns: Any, **kwargs: Any) -> pl.LazyFrame:
        events.append("select")
        values = columns[0] if len(columns) == 1 and isinstance(columns[0], list) else columns
        selected_columns.append([str(value) for value in values])
        return native_select(lazy_frame, *columns, **kwargs)

    def tracked_collect(lazy_frame: pl.LazyFrame, *args: Any, **kwargs: Any) -> pl.DataFrame:
        events.append("collect")
        result = cast(pl.DataFrame, native_collect(lazy_frame, *args, **kwargs))
        assert result.height <= 2
        assert result.columns == [row_id, "selected"]
        return result

    monkeypatch.setattr(pl.LazyFrame, "select", tracked_select)
    monkeypatch.setattr(pl.LazyFrame, "collect", tracked_collect)

    page = engine.page(
        frame,
        0,
        2,
        total_rows=source.height,
        column_projection=[(1, "stable:selected")],
    )

    if extension == "parquet":
        # Native text is prepared on the bounded resident frame after the source collect.
        assert events[:2] == ["select", "collect"]
        assert all(event == "collect" for event in events[2:])
    else:
        assert events == ["select", "collect"]
    assert selected_columns == [[row_id, "selected"]]
    assert page["columnIds"] == ["stable:selected"]
    assert [row["values"][0]["display"] for row in page["rows"]] == (
        ["1973-03-03T09:46:40.000000001", "1973-03-03T09:46:40.000000002"] if extension == "parquet" else ["30", "40"]
    )


def test_polars_column_values_and_parquet(tmp_path):
    frame = pl.DataFrame({"group": ["a", "a", "b"], "value": [1, 2, 3]})
    path = tmp_path / "sample.parquet"
    frame.write_parquet(path)

    manager = SessionManager()
    opened = manager.open_session({"kind": "file", "label": "sample.parquet", "path": str(path)}, backend="polars")
    values = manager.get_column_values(opened["metadata"]["sessionId"], 0, "group", {"filters": [], "sort": []})

    assert values["values"] == [
        {"value": "a", "count": 2, "selectionValue": typed_selection_value("a", "string")},
        {"value": "b", "count": 1, "selectionValue": typed_selection_value("b", "string")},
    ]


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_count_labels_keep_public_profiles_and_value_choices(
    lazy: bool, monkeypatch: pytest.MonkeyPatch
) -> None:
    import __main__

    values = ["b", "d", "a", "c", "a", "b", "a", None]
    frame = pl.DataFrame({"count": values, "count_": list(reversed(values))})
    before = frame.clone()
    source = frame.lazy() if lazy else frame
    monkeypatch.setattr(__main__, "count_label_source", source, raising=False)
    manager = SessionManager()
    view = {"filters": [], "sort": []}
    expected = [
        {"value": value, "count": count, "selectionValue": typed_selection_value(value, "string")}
        for value, count in (("a", 3), ("b", 2), ("c", 1), ("d", 1))
    ]
    try:
        opened = manager.open_session(
            {"kind": "notebookVariable", "label": "count_label_source", "variableName": "count_label_source"},
            backend="polars",
            page_size=8,
        )
        session_id = opened["metadata"]["sessionId"]
        schema = opened["metadata"]["schema"]
        assert [column["name"] for column in schema] == ["count", "count_"]
        for column in schema:
            for limit, has_more in ((3, True), (4, False)):
                choices = manager.get_column_values(session_id, 0, column["name"], view, limit=limit)
                assert choices == {
                    "kind": "columnValues",
                    "revision": 0,
                    "column": column["name"],
                    "values": expected[:limit],
                    "hasMore": has_more,
                }
            search = manager.get_column_values(session_id, 0, column["name"], view, search="B", limit=1)
            assert search["values"] == [expected[1]]
            assert search["hasMore"] is False
            summary = manager.get_summary(session_id, 0, view, [column["id"]])["summaries"][0]
            assert summary["columnId"] == column["id"]
            assert (summary["totalCount"], summary["nullCount"], summary["nanCount"], summary["distinctCount"]) == (
                8,
                1,
                0,
                4,
            )
            assert {item["value"]: item["count"] for item in summary["topValues"]} == {"a": 3, "b": 2, "c": 1, "d": 1}
        summaries = manager.get_summary(session_id, 0, view)["summaries"]
        assert [summary["columnId"] for summary in summaries] == [column["id"] for column in schema]
        page = manager.get_page(session_id, 0, 0, 8, view)
        assert page["revision"] == 0
        assert page["page"]["rows"] == opened["page"]["rows"]
        assert __main__.count_label_source is source
        assert (source.collect() if isinstance(source, pl.LazyFrame) else source).equals(before)
        assert frame.schema == before.schema
    finally:
        manager.close_all()


@pytest.mark.parametrize(
    ("column", "position"),
    [("__open_wrangler_count_0", 0), ("__open_wrangler_count_0", 1), ("__open_wrangler_count_1", 1)],
)
def test_lazy_polars_count_labels_keep_projected_and_full_profiles(
    column: str, position: int, monkeypatch: pytest.MonkeyPatch
) -> None:
    import __main__

    frame = pl.DataFrame({column: [True, False, True, None], "sibling": ["a", "b", "a", None]})
    if position == 1:
        frame = frame.select("sibling", column)
    source = frame.lazy()
    before = source.explain()
    monkeypatch.setattr(__main__, "projected_count_source", source, raising=False)
    manager = SessionManager()
    view = {"filters": [], "sort": []}
    try:
        opened = manager.open_session(
            {"kind": "notebookVariable", "label": "projected_count_source", "variableName": "projected_count_source"},
            backend="polars",
        )
        session_id = opened["metadata"]["sessionId"]
        schema = opened["metadata"]["schema"]
        column_id = next(item["id"] for item in schema if item["name"] == column)
        for projection in (None, [column_id]):
            response = manager.get_summary(session_id, 0, view, projection)
            assert response["revision"] == 0
            assert [item["columnId"] for item in response["summaries"]] == (
                [item["id"] for item in schema] if projection is None else [column_id]
            )
            summary = next(item for item in response["summaries"] if item["column"] == column)
            assert (summary["totalCount"], summary["nullCount"], summary["nanCount"], summary["distinctCount"]) == (
                4,
                1,
                0,
                2,
            )
            assert summary["topValues"] == [
                {"value": "True", "count": 2, "selectionValue": typed_selection_value(True, "boolean")},
                {"value": "False", "count": 1, "selectionValue": typed_selection_value(False, "boolean")},
            ]
            assert summary["visualization"] == {"kind": "boolean", "trueCount": 2, "falseCount": 1}
        assert __main__.projected_count_source is source
        assert source.explain() == before
        assert source.collect().equals(frame)
    finally:
        manager.close_all()


def test_polars_excel_session_refuses_missing_literal_path_without_reading_siblings(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from openpyxl import Workbook

    selected = tmp_path / "[a].xlsx"
    sibling = tmp_path / "a.xlsx"
    escaped = tmp_path / "[[]a].xlsx"
    for path, value in ((selected, 11), (sibling, 22), (escaped, 33)):
        workbook = Workbook()
        try:
            sheet = workbook.active
            assert sheet is not None
            sheet.title = "Sales"
            sheet.append(["value"])
            sheet.append([value])
            workbook.save(path)
        finally:
            workbook.close()
    before = {path: path.read_bytes() for path in (selected, sibling, escaped)}
    source = {
        "kind": "file",
        "label": selected.name,
        "path": str(selected),
        "importOptions": {"sheetName": "Sales"},
    }
    read_excel = pl.read_excel
    handles: list[io.BufferedReader] = []

    def track_read(handle: io.BufferedReader, **options: Any) -> pl.DataFrame:
        assert isinstance(handle, io.BufferedReader)
        assert not handle.closed
        handles.append(handle)
        return read_excel(handle, **options)

    monkeypatch.setattr(pl, "read_excel", track_read)
    manager = SessionManager()
    try:
        opened = manager.open_session(source, backend="polars")
        assert opened["metadata"]["source"] == source
        assert [row["values"][0]["display"] for row in opened["page"]["rows"]] == ["11"]
        assert len(handles) == 1 and handles[0].closed
        with pytest.raises(EngineError, match="missing"):
            manager.open_session({**source, "importOptions": {"sheetName": "missing"}}, backend="polars")
        assert len(handles) == 2 and handles[1].closed
        assert {path: path.read_bytes() for path in before} == before
        selected.unlink()
        with pytest.raises(EngineError, match="Could not read"):
            manager.open_session(source, backend="polars")
        assert not selected.exists()
        assert list(manager.sessions) == [opened["metadata"]["sessionId"]]
        page = manager.get_page(opened["metadata"]["sessionId"], 0, 0, 5, opened["metadata"]["filterModel"])
        assert [row["values"][0]["display"] for row in page["page"]["rows"]] == ["11"]
    finally:
        manager.close_all()
        assert manager.sessions == {}
        assert manager._opening_engines == {}
        assert all(handle.closed for handle in handles)
        if selected.exists():
            assert selected.read_bytes() == before[selected]
        assert sibling.read_bytes() == before[sibling]
        assert escaped.read_bytes() == before[escaped]


def test_polars_excel_reader_refuses_python_buffering_after_owned_path_removal(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "removed.xlsx"
    before = b"This source must never be read into a Python buffer."
    path.write_bytes(before)
    read_excel = pl.read_excel
    handles: list[io.BufferedReader] = []

    def remove_before_parse(handle: io.BufferedReader, **options: Any) -> pl.DataFrame:
        assert isinstance(handle, io.BufferedReader)
        handles.append(handle)
        try:
            path.unlink()
        except PermissionError:
            if os.name == "nt":
                pytest.skip("Windows refused removal of the owned open workbook handle.")
            raise
        return read_excel(handle, **options)

    monkeypatch.setattr(pl, "read_excel", remove_before_parse)
    engine = PolarsEngine()
    try:
        with pytest.raises(EngineError, match="without buffering the whole workbook"):
            engine.read_file(str(path))
        assert not path.exists()
    finally:
        engine.close()
        assert len(handles) == 1 and handles[0].closed
        if path.exists():
            assert path.read_bytes() == before


def test_polars_excel_reader_pins_the_probed_calamine_engine(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[tuple[str, dict[str, object]]] = []
    handles: list[io.BufferedReader] = []

    def read_excel(source: io.BufferedReader, **options: object) -> pl.DataFrame:
        assert isinstance(source, io.BufferedReader)
        assert not source.closed
        handles.append(source)
        calls.append((Path(source.name).name, options))
        return pl.DataFrame({"value": [1]})

    monkeypatch.setattr(pl, "read_excel", read_excel)
    runtime = PolarsEngine()

    for name in ("default.xlsx", "modern.xlsx", "legacy.xls"):
        (tmp_path / name).write_bytes(b"Native reader driver fixture")
    try:
        runtime.read_file(str(tmp_path / "default.xlsx"))
        runtime.read_file(str(tmp_path / "modern.xlsx"), {"sheetIndex": 1})
        runtime.read_file(str(tmp_path / "legacy.xls"), {"sheetName": " résumé "})
    finally:
        runtime.close()
        assert all(handle.closed for handle in handles)

    assert calls == [
        ("default.xlsx", {"sheet_id": 1, "engine": "calamine"}),
        ("modern.xlsx", {"sheet_id": 2, "engine": "calamine"}),
        ("legacy.xls", {"sheet_name": " résumé ", "engine": "calamine"}),
    ]


def test_lazy_polars_schema_discovery_does_not_collect_column_profiles(monkeypatch):
    frame = pl.DataFrame({"complete": [1, 2], "with_null": [1, None]}).lazy()

    def reject_collect(*_args, **_kwargs):
        raise AssertionError("Lazy schema discovery must not execute the data plan")

    monkeypatch.setattr(pl.LazyFrame, "collect", reject_collect)

    schema = PolarsEngine().schema(frame)

    assert [(column["name"], column["rawType"], column["nullable"]) for column in schema] == [
        ("complete", "Int64", True),
        ("with_null", "Int64", True),
    ]


def test_polars_normalize_keeps_lazyframes_and_converts_series() -> None:
    engine = PolarsEngine()
    lazy = pl.DataFrame({"value": [1, 2]}).lazy()

    assert engine.normalize(lazy) is lazy
    normalized_series = engine.normalize(pl.Series("value", [1, 2]))
    assert isinstance(normalized_series, pl.DataFrame)
    assert normalized_series.to_dict(as_series=False) == {"value": [1, 2]}


@pytest.mark.parametrize("format_name", ["csv", "parquet"])
def test_lazy_polars_export_streams_to_the_exact_reserved_file_object(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    format_name: Literal["csv", "parquet"],
) -> None:
    frame = pl.LazyFrame({"value": [1, 2, 3]})
    destination = tmp_path / f"lazy-stream.{format_name}"
    destination.touch()
    identity = _regular_file_identity(destination)
    writer_path = ExportWriterPath(destination, *identity)
    observed_writers: list[Any] = []
    native_sink = pl.LazyFrame.sink_csv if format_name == "csv" else pl.LazyFrame.sink_parquet

    def observed_sink(lazy_frame: pl.LazyFrame, writer: Any, *args: Any, **kwargs: Any) -> Any:
        assert not isinstance(writer, (str, Path))
        assert callable(getattr(writer, "write", None))
        observed_writers.append(writer)
        return native_sink(lazy_frame, writer, *args, **kwargs)

    monkeypatch.setattr(pl.LazyFrame, f"sink_{format_name}", observed_sink)

    monkeypatch.setattr(
        pl.DataFrame,
        "to_pandas",
        lambda *_args, **_kwargs: pytest.fail("Polars export must not convert to Pandas"),
        raising=False,
    )

    PolarsEngine().export_data(frame, writer_path, export_options(format_name))

    assert len(observed_writers) == 1
    assert observed_writers[0].closed is True
    assert _regular_file_identity(destination) == identity
    result = pl.read_csv(destination) if format_name == "csv" else pl.read_parquet(destination)
    assert result.to_dict(as_series=False) == {"value": [1, 2, 3]}


def test_live_notebook_lazyframe_stays_lazy_through_bounded_queries_edit_export_and_close(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import __main__

    row_count = 1_000_003
    source_path = tmp_path / "large-live-source.parquet"
    (
        pl.LazyFrame()
        .select(pl.int_range(0, row_count).alias("row"))
        .with_columns(
            [
                (pl.col("row") % 7).alias("group"),
                (pl.lit("ROW-") + (pl.col("row") % 17).cast(pl.String)).alias("label"),
                ((pl.col("row") * 13) % 10_000).alias("score"),
            ]
        )
        .sink_parquet(source_path)
    )
    live = pl.scan_parquet(source_path)
    original_schema = live.collect_schema()
    original_plan = live.explain(optimized=False)
    monkeypatch.setattr(__main__, "large_live_lazyframe", live, raising=False)

    native_collect = pl.LazyFrame.collect
    native_collect_all = pl.collect_all
    native_to_list = pl.Series.to_list
    collected_heights: list[int] = []
    to_list_lengths: list[int] = []
    captured_heights: list[int] = []

    def bounded_collect(frame: pl.LazyFrame, *args: Any, **kwargs: Any) -> pl.DataFrame:
        result = cast(pl.DataFrame, native_collect(frame, *args, **kwargs))
        collected_heights.append(result.height)
        assert result.height <= 20, "A live LazyFrame query collected an unbounded result."
        return result

    def bounded_collect_all(frames: Any, *args: Any, **kwargs: Any) -> list[pl.DataFrame]:
        results = native_collect_all(frames, *args, **kwargs)
        if len(frames) == 1 and frames[0] is live:
            assert kwargs == {"engine": "in-memory"}
            captured_heights.extend(result.height for result in results)
            return results
        collected_heights.extend(result.height for result in results)
        assert all(result.height <= 20 for result in results), "A live LazyFrame profile collected unbounded results."
        return results

    def bounded_to_list(series: pl.Series) -> list[Any]:
        to_list_lengths.append(len(series))
        assert len(series) <= 20, "A live LazyFrame query converted an unbounded series to a list."
        return native_to_list(series)

    def reject_to_pandas(*_args: Any, **_kwargs: Any) -> None:
        raise AssertionError("A live Polars LazyFrame must remain Polars-native.")

    monkeypatch.setattr(pl.LazyFrame, "collect", bounded_collect)
    monkeypatch.setattr(pl, "collect_all", bounded_collect_all)
    monkeypatch.setattr(pl.Series, "to_list", bounded_to_list)
    monkeypatch.setattr(pl.DataFrame, "to_pandas", reject_to_pandas, raising=False)

    manager = SessionManager()
    opened = manager.open_session(
        {
            "kind": "notebookVariable",
            "label": "large_live_lazyframe",
            "variableName": "large_live_lazyframe",
        },
        backend="polars",
        mode="editing",
        page_size=3,
        column_offset=2,
        column_limit=1,
    )
    session_id = opened["metadata"]["sessionId"]
    runtime = manager.sessions[session_id]

    assert opened["metadata"]["shape"] == {"rows": row_count, "columns": 4}
    assert opened["metadata"]["capabilities"]["lazy"] is True
    assert opened["summaries"] == []
    assert [column["name"] for column in opened["metadata"]["schema"]] == ["row", "group", "label", "score"]
    assert opened["page"]["columnIds"] == [opened["metadata"]["schema"][2]["id"]]
    assert [row["values"][0]["display"] for row in opened["page"]["rows"]] == ["ROW-0", "ROW-1", "ROW-2"]
    assert isinstance(runtime.original, pl.LazyFrame)
    assert isinstance(runtime.committed, pl.LazyFrame)

    filter_model = {
        "logic": "and",
        "filters": [
            {
                "column": "group",
                "type": "integer",
                "logic": "and",
                "valueFilter": None,
                "predicates": [{"kind": "predicate", "operator": "equals", "value": 3}],
            }
        ],
        "sort": [{"column": "score", "direction": "desc", "nulls": "last"}],
    }
    filtered_page = manager.get_page(
        session_id,
        0,
        120,
        5,
        filter_model,
        column_offset=0,
        column_limit=2,
    )
    assert filtered_page["metadata"]["capabilities"]["lazy"] is True
    assert filtered_page["page"]["columnIds"] == [
        opened["metadata"]["schema"][0]["id"],
        opened["metadata"]["schema"][1]["id"],
    ]
    assert len(filtered_page["page"]["rows"]) == 5
    assert all(row["values"][1]["display"] == "3" for row in filtered_page["page"]["rows"])

    score_id = opened["metadata"]["schema"][3]["id"]
    summary = manager.get_summary(session_id, 0, filter_model, [score_id])
    assert summary["summaries"][0]["totalCount"] == filtered_page["metadata"]["filteredShape"]["rows"]
    assert summary["summaries"][0]["numeric"]["max"] == 9999.0
    stats = manager.get_dataset_stats(session_id, 0, filter_model)
    assert stats["stats"]["missingCells"] == 0
    assert len(stats["stats"]["missingValuesByColumn"]) == 4

    label_column = opened["metadata"]["schema"][2]
    preview = manager.preview_step(
        session_id,
        0,
        {
            "id": "lower-live-label",
            "kind": "lowerText",
            "params": {
                "column": {"id": label_column["id"], "name": label_column["name"]},
                "newColumn": "label_lower",
            },
        },
        0,
        4,
        column_offset=3,
        column_limit=2,
    )
    assert preview["metadata"]["capabilities"]["lazy"] is True
    assert isinstance(runtime.draft_frame, pl.LazyFrame)
    assert preview["diff"]["addedColumns"] == ["label_lower"]
    assert preview["page"]["rows"][0]["values"][1]["display"].startswith("row-")

    applied = manager.apply_draft(session_id, 1, 0, 4, column_offset=3, column_limit=2)
    assert applied["metadata"]["capabilities"]["lazy"] is True
    assert isinstance(runtime.committed, pl.LazyFrame)

    export_path = tmp_path / "large-live-cleaned.parquet"
    exported = manager.export_data(
        session_id,
        2,
        str(export_path),
        export_options("parquet"),
        reserve_export_target(export_path),
    )
    assert exported["shape"] == {"rows": row_count, "columns": 5}
    exported_metrics = (
        pl.scan_parquet(export_path)
        .select(
            [
                pl.len().alias("rows"),
                pl.col("label_lower").first().alias("first_label"),
            ]
        )
        .collect(engine="streaming")
    )
    assert exported_metrics.row(0) == (row_count, "row-0")

    assert manager.close_session(session_id, 2) == {"kind": "sessionClosed", "sessionId": session_id}
    assert session_id not in manager.sessions
    assert live.collect_schema() == original_schema
    assert live.explain(optimized=False) == original_plan
    assert captured_heights == [row_count]
    assert collected_heights
    assert max(collected_heights) <= 20
    assert all(length <= 20 for length in to_list_lengths)


@pytest.mark.parametrize("kind", ["oneHotEncode", "multiLabelBinarize"])
def test_live_notebook_lazyframe_dynamic_encoder_returns_eager_preview(
    kind: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import __main__

    source_path = tmp_path / f"explicit-{kind}.parquet"
    pl.DataFrame({"category": ["a", "b", "a"], "tags": ["x|y", "y", "x"]}).write_parquet(source_path)
    live = pl.scan_parquet(source_path)
    monkeypatch.setattr(__main__, "explicit_lazy_encoder", live, raising=False)
    native_collect = pl.LazyFrame.collect
    collected_heights: list[int] = []

    def tracked_collect(frame: pl.LazyFrame, *args: Any, **kwargs: Any) -> pl.DataFrame:
        result = cast(pl.DataFrame, native_collect(frame, *args, **kwargs))
        collected_heights.append(result.height)
        return result

    monkeypatch.setattr(pl.LazyFrame, "collect", tracked_collect)
    manager = SessionManager()
    opened = manager.open_session(
        {
            "kind": "notebookVariable",
            "label": "explicit_lazy_encoder",
            "variableName": "explicit_lazy_encoder",
        },
        backend="polars",
        mode="editing",
        page_size=2,
    )
    session_id = opened["metadata"]["sessionId"]
    assert opened["metadata"]["capabilities"]["lazy"] is True
    assert max(collected_heights) <= 2
    collected_heights.clear()

    schema = {column["name"]: column for column in opened["metadata"]["schema"]}
    if kind == "oneHotEncode":
        params = {
            "columns": [{"id": schema["category"]["id"], "name": "category"}],
            "dropOriginal": False,
        }
    else:
        params = {
            "column": {"id": schema["tags"]["id"], "name": "tags"},
            "delimiter": "|",
            "dropOriginal": False,
        }
    preview = manager.preview_step(
        session_id,
        0,
        {"id": f"explicit-{kind}", "kind": kind, "params": params},
        0,
        2,
    )

    assert collected_heights[0] == 3
    assert max(collected_heights) == 3
    assert isinstance(manager.sessions[session_id].draft_frame, pl.DataFrame)
    assert preview["metadata"]["capabilities"]["lazy"] is False
    manager.close_session(session_id, 1)


def test_lazy_polars_numeric_summary_is_exact_with_only_bounded_collections(monkeypatch):
    row_count = 12_305
    values = pl.concat(
        [
            pl.int_range(0, row_count - 1, eager=True) % 101,
            pl.Series([1_000_000]),
        ]
    )
    source = pl.DataFrame({"value": pl.concat([pl.Series([-1] * 500), values])}).lazy()
    frame = PolarsEngine().apply_filter_model(
        source,
        {
            "logic": "and",
            "filters": [
                {
                    "column": "value",
                    "type": "integer",
                    "logic": "and",
                    "predicates": [{"operator": "gte", "value": 0}],
                }
            ],
            "sort": [],
        },
    )
    eager = values.cast(pl.Float64)
    collected_heights: list[int] = []
    original_collect = pl.LazyFrame.collect
    original_to_list = pl.Series.to_list

    def bounded_collect(lazy_frame: pl.LazyFrame, *args: Any, **kwargs: Any) -> pl.DataFrame:
        result = cast(pl.DataFrame, original_collect(lazy_frame, *args, **kwargs))
        assert isinstance(result, pl.DataFrame)
        collected_heights.append(result.height)
        assert result.height <= 20
        return result

    monkeypatch.setattr(pl.LazyFrame, "collect", bounded_collect)

    def bounded_to_list(series):
        assert len(series) <= 20
        return original_to_list(series)

    monkeypatch.setattr(pl.Series, "to_list", bounded_to_list)

    summary = PolarsEngine().summaries(frame, [(0, "c:value")])[0]

    assert collected_heights
    assert max(collected_heights) <= 20
    assert summary["totalCount"] == row_count
    assert summary["nullCount"] == 0
    assert summary["nanCount"] == 0
    assert summary["distinctCount"] == 102
    assert summary["numeric"] == {
        "min": 0.0,
        "max": 1_000_000.0,
        "mean": pytest.approx(eager.mean()),
        "median": pytest.approx(eager.median()),
        "std": pytest.approx(eager.std()),
        "sum": pytest.approx(eager.sum()),
        "exactSum": {
            "kind": "integer",
            "raw": 1_614_453,
            "display": "1614453",
            "isNull": False,
            "isNaN": False,
        },
        "exactMin": {
            "kind": "integer",
            "raw": 0,
            "display": "0",
            "isNull": False,
            "isNaN": False,
        },
        "exactMax": {
            "kind": "integer",
            "raw": 1_000_000,
            "display": "1000000",
            "isNull": False,
            "isNaN": False,
        },
    }
    visualization = summary["visualization"]
    assert "sampled" not in visualization
    assert len(visualization["bins"]) == 20
    assert visualization["bins"][0]["min"] == summary["numeric"]["min"]
    assert visualization["bins"][-1]["max"] == summary["numeric"]["max"]
    assert visualization["bins"][-1]["count"] == 1
    assert sum(bin_["count"] for bin_ in visualization["bins"]) == row_count


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_profiles_a_wide_int64_projection_with_exact_native_sums(
    lazy: bool,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    column_count = 64
    row_count = 257
    frame = pl.DataFrame(
        {
            f"value_{column}": pl.Series(
                [(row % 17) + column for row in range(row_count)],
                dtype=pl.Int64,
            )
            for column in range(column_count)
        }
    )
    source = frame.lazy() if lazy else frame
    native_collect = pl.LazyFrame.collect
    native_collect_all = pl.collect_all
    collected_heights: list[int] = []

    def bounded_collect(lazy_frame: pl.LazyFrame, *args: Any, **kwargs: Any) -> pl.DataFrame:
        result = cast(pl.DataFrame, native_collect(lazy_frame, *args, **kwargs))
        collected_heights.append(result.height)
        assert result.height <= 20
        return result

    def bounded_collect_all(frames: Any, *args: Any, **kwargs: Any) -> list[pl.DataFrame]:
        results = native_collect_all(frames, *args, **kwargs)
        collected_heights.extend(result.height for result in results)
        assert all(result.height <= 20 for result in results)
        return results

    if lazy:
        monkeypatch.setattr(pl.LazyFrame, "collect", bounded_collect)
        monkeypatch.setattr(pl, "collect_all", bounded_collect_all)
    projection = [(position, f"c:{position}") for position in range(column_count)]

    summaries = PolarsEngine().summaries(source, projection)

    base_sum = sum(row % 17 for row in range(row_count))
    assert [summary["columnId"] for summary in summaries] == [column_id for _, column_id in projection]
    assert [summary["numeric"]["exactSum"]["raw"] for summary in summaries] == [
        base_sum + (position * row_count) for position in range(column_count)
    ]
    if lazy:
        assert collected_heights
        assert max(collected_heights) <= 20


def test_live_lazy_polars_profiles_many_horizontal_column_windows_after_capture(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import __main__

    row_count = 10_003
    column_count = 40
    window_size = 8
    source_path = tmp_path / "wide-live-profile.parquet"
    rows = pl.int_range(0, row_count, eager=True)
    pl.DataFrame(
        {f"value_{column}": ((rows % 17) + column).cast(pl.Int64) for column in range(column_count)}
    ).write_parquet(source_path)
    live = pl.scan_parquet(source_path)
    original_schema = live.collect_schema()
    original_plan = live.explain(optimized=False)
    monkeypatch.setattr(__main__, "wide_live_profile", live, raising=False)

    native_collect = pl.LazyFrame.collect
    native_collect_all = pl.collect_all
    native_to_list = pl.Series.to_list
    collected_heights: list[int] = []
    to_list_lengths: list[int] = []
    captured_heights: list[int] = []

    def bounded_collect(lazy_frame: pl.LazyFrame, *args: Any, **kwargs: Any) -> pl.DataFrame:
        result = cast(pl.DataFrame, native_collect(lazy_frame, *args, **kwargs))
        collected_heights.append(result.height)
        assert result.height <= 20, "A horizontal LazyFrame profile collected an unbounded result."
        return result

    def bounded_collect_all(frames: Any, *args: Any, **kwargs: Any) -> list[pl.DataFrame]:
        results = native_collect_all(frames, *args, **kwargs)
        if len(frames) == 1 and frames[0] is live:
            assert kwargs == {"engine": "in-memory"}
            captured_heights.extend(result.height for result in results)
            return results
        collected_heights.extend(result.height for result in results)
        assert all(result.height <= 20 for result in results), "A horizontal profile collected unbounded results."
        return results

    def bounded_to_list(series: pl.Series) -> list[Any]:
        to_list_lengths.append(len(series))
        assert len(series) <= 20, "A horizontal profile converted an unbounded Series to a list."
        return native_to_list(series)

    monkeypatch.setattr(pl.LazyFrame, "collect", bounded_collect)
    monkeypatch.setattr(pl, "collect_all", bounded_collect_all)
    monkeypatch.setattr(pl.Series, "to_list", bounded_to_list)

    manager = SessionManager()
    opened = manager.open_session(
        {
            "kind": "notebookVariable",
            "label": "wide_live_profile",
            "variableName": "wide_live_profile",
        },
        backend="polars",
        mode="editing",
        page_size=2,
        column_limit=window_size,
    )
    session_id = opened["metadata"]["sessionId"]
    schema = opened["metadata"]["schema"]
    empty_view = {"filters": [], "sort": []}
    profiled_ids: list[str] = []
    pages = [opened]
    for column_offset in range(window_size, column_count, window_size):
        pages.append(
            manager.get_page(
                session_id,
                0,
                0,
                2,
                empty_view,
                column_offset=column_offset,
                column_limit=window_size,
            )
        )

    base_sum = sum(row % 17 for row in range(row_count))
    for page_response in pages:
        for column_id in page_response["page"]["columnIds"]:
            position = next(column["position"] for column in schema if column["id"] == column_id)
            response = manager.get_summary(session_id, 0, empty_view, [column_id])
            summary = response["summaries"][0]
            assert summary["numeric"]["exactSum"]["raw"] == base_sum + (position * row_count)
            profiled_ids.append(column_id)

    assert opened["metadata"]["capabilities"]["lazy"] is True
    assert isinstance(manager.sessions[session_id].original, pl.LazyFrame)
    assert profiled_ids == [column["id"] for column in schema]
    assert manager.close_session(session_id, 0) == {"kind": "sessionClosed", "sessionId": session_id}
    assert live.collect_schema() == original_schema
    assert live.explain(optimized=False) == original_plan
    assert captured_heights == [row_count]
    assert collected_heights and max(collected_heights) <= 20
    assert all(length <= 20 for length in to_list_lengths)


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_numeric_summaries_publish_lossless_wide_extrema_and_sums(
    lazy: bool,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def reject_conversion(*_args: Any, **_kwargs: Any) -> None:
        raise AssertionError("Polars summaries must remain native.")

    monkeypatch.setattr(pl.DataFrame, "to_pandas", reject_conversion, raising=False)
    wide_values = [-(10**30) + 2, 10**30 + 3, 10**30 + 1]
    unsigned_values = [10**30 + 5, 7, None]
    decimal_values = [
        Decimal("-12345678901234567890.123456789012345678"),
        Decimal("98765432109876543210.987654321098765432"),
        None,
    ]
    frame = pl.DataFrame(
        {
            "wide": pl.Series(wide_values, dtype=pl.Int128),
            "unsigned": pl.Series(unsigned_values, dtype=pl.UInt128),
            "amount": pl.Series(decimal_values, dtype=pl.Decimal(38, 18)),
        }
    )

    summaries = PolarsEngine().summaries(frame.lazy() if lazy else frame)

    assert summaries[0]["numeric"]["exactMin"] == {
        "kind": "integer",
        "raw": str(min(wide_values)),
        "display": str(min(wide_values)),
        "isNull": False,
        "isNaN": False,
    }
    assert summaries[0]["numeric"]["exactMax"] == {
        "kind": "integer",
        "raw": str(max(wide_values)),
        "display": str(max(wide_values)),
        "isNull": False,
        "isNaN": False,
    }
    assert summaries[0]["numeric"]["exactSum"]["display"] == str(sum(wide_values))
    assert summaries[1]["numeric"]["exactSum"]["display"] == str(sum(value or 0 for value in unsigned_values))
    assert summaries[2]["numeric"]["exactMin"]["display"] == str(decimal_values[0])
    assert summaries[2]["numeric"]["exactMax"]["display"] == str(decimal_values[1])
    assert summaries[2]["numeric"]["exactSum"]["display"] == "86419753208641975320.864197532086419754"
    assert summaries[2]["numeric"]["exactMin"]["kind"] == "decimal"
    assert summaries[2]["numeric"]["exactMax"]["kind"] == "decimal"


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_summary_omits_non_finite_statistics_but_keeps_finite_histogram_values(lazy: bool):
    frame = pl.DataFrame({"value": [1.0, float("inf")]})
    summary = PolarsEngine().summaries(frame.lazy() if lazy else frame)[0]

    assert summary["numeric"] == {"min": 1.0}
    assert summary["visualization"] == {"kind": "numeric", "bins": [{"min": 1.0, "max": 1.0, "count": 1}]}


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_numeric_histogram_counts_all_valid_values_after_nulls(lazy: bool):
    row_count = 16_384
    values = [None if index % 2 == 0 else float(index) for index in range(row_count)]
    frame = pl.DataFrame({"value": values})
    source = frame.lazy() if lazy else frame

    first = PolarsEngine().summaries(source, [(0, "c:value")])[0]
    second = PolarsEngine().summaries(source, [(0, "c:value")])[0]

    assert first["numeric"]["min"] == 1.0
    assert first["numeric"]["max"] == float(row_count - 1)
    assert first["distinctCount"] == row_count // 2
    assert first["visualization"] == second["visualization"]
    assert first["visualization"]["bins"]
    histogram_count = sum(bin_["count"] for bin_ in first["visualization"]["bins"])
    assert histogram_count == row_count // 2
    assert "sampled" not in first["visualization"]


def test_polars_eager_numeric_histogram_preserves_boundary_and_rounded_edge_assignment() -> None:
    adjacent_extremes = [1e308]
    for _index in range(24):
        adjacent_extremes.append(nextafter(adjacent_extremes[-1], float("inf")))

    for values in ([float(value) for value in range(21)], adjacent_extremes):
        summary = PolarsEngine().summaries(pl.DataFrame({"value": values}))[0]

        assert summary["visualization"] == engine_base.numeric_visualization(values)


def test_polars_eager_numeric_summaries_never_materialize_python_value_lists(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def reject_list_materialization(*_args: Any, **_kwargs: Any) -> None:
        raise AssertionError("Eager Polars numeric summaries must remain in native expressions.")

    monkeypatch.setattr(pl.Series, "to_list", reject_list_materialization)
    monkeypatch.setattr(engine_base, "numeric_visualization", reject_list_materialization)
    monkeypatch.setattr(polars_engine, "numeric_visualization", reject_list_materialization, raising=False)
    frame = pl.DataFrame(
        {
            "integer": pl.Series([None, *[value % 101 for value in range(4_096)]], dtype=pl.Int64),
            "floating": pl.Series(
                [None, float("nan"), float("-inf"), float("inf"), *[float(value % 101) for value in range(4_093)]],
                dtype=pl.Float64,
            ),
        }
    )

    summaries = {summary["column"]: summary for summary in PolarsEngine().summaries(frame)}

    integer = summaries["integer"]
    assert (integer["totalCount"], integer["nullCount"], integer["nanCount"]) == (4_097, 1, 0)
    assert integer["distinctCount"] == 101
    assert integer["topValues"][0] == {"value": "0", "count": 41, "selectionValue": typed_selection_value(0, "integer")}
    assert integer["numeric"]["exactMin"]["display"] == "0"
    assert integer["numeric"]["exactMax"]["display"] == "100"
    assert len(integer["visualization"]["bins"]) == 20
    assert sum(bin_["count"] for bin_ in integer["visualization"]["bins"]) == 4_096
    floating = summaries["floating"]
    assert (floating["totalCount"], floating["nullCount"], floating["nanCount"]) == (4_097, 1, 1)
    assert floating["topValues"][0] == {
        "value": "0.0",
        "count": 41,
        "selectionValue": typed_selection_value(0.0, "float"),
    }
    assert sum(bin_["count"] for bin_ in floating["visualization"]["bins"]) == 4_093


def test_polars_eager_boolean_summaries_reuse_native_counts_without_materializing_lists(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def reject_list_materialization(*_args: Any, **_kwargs: Any) -> None:
        raise AssertionError("Eager Polars boolean summaries must reuse native value counts.")

    monkeypatch.setattr(pl.Series, "to_list", reject_list_materialization)
    monkeypatch.setattr(engine_base, "boolean_visualization", reject_list_materialization)
    monkeypatch.setattr(polars_engine, "boolean_visualization", reject_list_materialization, raising=False)
    frame = pl.DataFrame(
        {
            "count": pl.Series([True, False, True, True, False, True], dtype=pl.Boolean),
            "count_": pl.Series([True, False, True, None, True, None], dtype=pl.Boolean),
            "all_null": pl.Series([None] * 6, dtype=pl.Boolean),
        }
    )

    summaries = {summary["column"]: summary for summary in PolarsEngine().summaries(frame)}

    expected = {
        "count": (0, 2, {"kind": "boolean", "trueCount": 4, "falseCount": 2}),
        "count_": (2, 2, {"kind": "boolean", "trueCount": 3, "falseCount": 1}),
        "all_null": (6, 0, {"kind": "boolean", "trueCount": 0, "falseCount": 0}),
    }
    for column, (null_count, distinct_count, visualization) in expected.items():
        summary = summaries[column]
        assert summary["type"] == "boolean"
        assert summary["totalCount"] == 6
        assert summary["nullCount"] == null_count
        assert summary["nanCount"] == 0
        assert summary["distinctCount"] == distinct_count
        assert {item["value"]: item["count"] for item in summary["topValues"]} == {
            str(value): count
            for value, count in ((True, visualization["trueCount"]), (False, visualization["falseCount"]))
            if count
        }
        assert summary["visualization"] == visualization


def test_lazy_polars_header_stats_collect_only_scalar_results(monkeypatch):
    frame = pl.DataFrame(
        {
            "value": [1.0, 1.0, None, float("nan"), float("nan")],
            "group": ["x", "x", "y", "z", "z"],
        }
    ).lazy()
    collected_shapes: list[tuple[int, int]] = []
    original_collect_all = pl.collect_all

    def scalar_collect_all(queries, *args: Any, **kwargs: Any):
        results = original_collect_all(queries, *args, **kwargs)
        collected_shapes.extend(result.shape for result in results)
        assert all(result.height <= 1 for result in results)
        return results

    monkeypatch.setattr(pl, "collect_all", scalar_collect_all)

    stats = PolarsEngine().header_stats(frame)

    assert collected_shapes == [(1, 4), (1, 1)]
    assert stats == {
        "missingCells": 3,
        "missingRows": 3,
        "duplicateRows": 2,
        "missingValuesByColumn": [
            {"column": "value", "count": 3},
            {"column": "group", "count": 0},
        ],
    }


@pytest.mark.parametrize("affinity", [None, "streaming"])
@pytest.mark.parametrize("case", ["mixed", "single", "empty", "all-null", "no-visible"])
def test_lazy_polars_object_stats_keep_exact_missing_metrics_without_grouping(
    affinity: Literal["streaming"] | None, case: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = pl.DataFrame(
        {
            "value": pl.Series(["kept", None, "kept", "nan"], dtype=pl.Object),
            "number": [-0.0, None, -0.0, float("nan")],
        }
    )
    if case == "single":
        source = source.select("value")
    elif case == "empty":
        source = source.head(0)
    elif case == "all-null":
        source = source.slice(1, 1)
    elif case == "no-visible":
        source = source.select([])
    before = source.clone()
    counts = {"mixed": [1, 2], "single": [1], "empty": [0, 0], "all-null": [1, 1], "no-visible": []}[case]
    expected = {
        "missingCells": sum(counts),
        "missingRows": {"mixed": 2, "single": 1, "empty": 0, "all-null": 1, "no-visible": 0}[case],
        "duplicateRows": None if source.height else 0,
        "missingValuesByColumn": [
            {"column": column, "count": count} for column, count in zip(source.columns, counts, strict=True)
        ],
    }
    engine = PolarsEngine()
    with pl.Config(set_engine_affinity=None):
        assert engine.header_stats(source) == {
            **expected,
            "duplicateRows": 1 if case in {"mixed", "single"} else 0,
        }

    collected_shapes: list[tuple[int, int]] = []
    native_collect_all = pl.collect_all

    def scalar_collect_all(queries: Any, **kwargs: Any) -> list[pl.DataFrame]:
        assert kwargs == {"engine": "streaming"}
        assert len(queries) == 1
        outputs = native_collect_all(queries, **kwargs)
        collected_shapes.extend(output.shape for output in outputs)
        assert all(output.height == 1 for output in outputs)
        return outputs

    def reject_unique(*_args: Any, **_kwargs: Any) -> Any:
        raise AssertionError("Lazy Object statistics must not construct a duplicate grouping query.")

    monkeypatch.setattr(pl, "collect_all", scalar_collect_all)
    monkeypatch.setattr(pl.LazyFrame, "unique", reject_unique)
    with pl.Config(set_engine_affinity=affinity):
        assert engine.header_stats(source.lazy()) == expected
    assert collected_shapes == ([(1, 2 + source.width)] if source.width else [])
    assert source.schema == before.schema
    if "value" in source.columns:
        assert source["value"].to_list() == before["value"].to_list()
        assert source.drop("value").equals(before.drop("value"))


def test_live_lazy_object_dataset_stats_preserve_filtered_page_and_source(monkeypatch: pytest.MonkeyPatch) -> None:
    import __main__

    source = pl.DataFrame(
        {
            "value": pl.Series(["kept", None, "kept", "outside"], dtype=pl.Object),
            "number": [1.0, float("nan"), 1.0, None],
            "keep": [True, True, True, False],
        }
    )
    before = source.clone()
    lazy = source.lazy()
    monkeypatch.setattr(__main__, "object_stats_source", lazy, raising=False)
    manager = SessionManager()
    try:
        opened = manager.open_session(
            {"kind": "notebookVariable", "variableName": "object_stats_source"},
            backend="polars",
            page_size=4,
            column_offset=1,
            column_limit=1,
        )
        sid = opened["metadata"]["sessionId"]
        view = {
            "filters": [{"column": "keep", "type": "boolean", "predicates": [{"operator": "equals", "value": True}]}],
            "sort": [{"column": "number", "direction": "desc", "nulls": "last"}],
        }
        page = manager.get_page(sid, 0, 0, 4, view, column_offset=1, column_limit=1)
        assert page["metadata"]["filteredShape"] == {"rows": 3, "columns": 3}
        assert page["page"]["columnIds"] == [opened["metadata"]["schema"][1]["id"]]
        stats = manager.get_dataset_stats(sid, 0, view)
        assert stats == {
            "kind": "datasetStats",
            "revision": 0,
            "stats": {
                "missingCells": 2,
                "missingRows": 1,
                "duplicateRows": None,
                "missingValuesByColumn": [
                    {"column": "value", "count": 1},
                    {"column": "number", "count": 1},
                    {"column": "keep", "count": 0},
                ],
            },
        }
        assert json.loads(json.dumps(stats, allow_nan=False)) == stats
        assert manager.get_page(sid, 0, 0, 4, view, column_offset=1, column_limit=1) == page
        assert manager.sessions[sid].plan == []
        assert isinstance(manager.sessions[sid].committed, pl.LazyFrame)
        assert __main__.object_stats_source is lazy
        assert source.schema == before.schema
        assert source["value"].to_list() == before["value"].to_list()
        assert source.drop("value").equals(before.drop("value"))
    finally:
        manager.close_all()


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_summary_excludes_null_and_nan_from_values_and_numeric_metrics(lazy: bool):
    frame = pl.DataFrame({"value": [1.0, None, float("nan"), 1.0]})
    source = frame.lazy() if lazy else frame

    summary = PolarsEngine().summaries(source, [(0, "c:value")])[0]

    assert summary["nullCount"] == 1
    assert summary["nanCount"] == 1
    assert summary["distinctCount"] == 1
    assert summary["topValues"] == [{"value": "1.0", "count": 2, "selectionValue": typed_selection_value(1.0, "float")}]
    assert summary["numeric"] == {
        "min": 1.0,
        "max": 1.0,
        "mean": 1.0,
        "median": 1.0,
        "std": 0.0,
        "sum": 2.0,
    }


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_text_summaries_are_exact_for_unicode_empty_and_all_null_without_pandas(
    lazy: bool,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail_to_pandas(*_args: Any, **_kwargs: Any) -> None:
        raise AssertionError("Polars text summaries must not convert to Pandas")

    monkeypatch.setattr(pl.DataFrame, "to_pandas", fail_to_pandas, raising=False)
    frame = pl.DataFrame(
        {
            "text": pl.Series([None, "", "A", "é", "e\u0301", "😀"], dtype=pl.String),
            "all_null": pl.Series([None, None, None, None, None, None], dtype=pl.String),
        }
    )

    summaries = {summary["column"]: summary for summary in PolarsEngine().summaries(frame.lazy() if lazy else frame)}

    assert summaries["text"]["text"] == {
        "emptyCount": 1,
        "minLength": 0,
        "maxLength": 2,
        "meanLength": 1.0,
    }
    assert summaries["all_null"]["text"] == {"emptyCount": 0}


def test_lazy_polars_nested_summaries_keep_exact_display_counts():
    frame = pl.DataFrame(
        {
            "items": [[1, 2], [1, 2], [3], None],
            "record": [{"x": 1}, {"x": 1}, {"x": 2}, None],
        }
    ).lazy()

    summaries = PolarsEngine().summaries(frame)

    assert summaries[0]["distinctCount"] == 2
    assert summaries[0]["topValues"] == [
        {"value": "[1,2]", "count": 2, "selectionValue": None},
        {"value": "[3]", "count": 1, "selectionValue": None},
    ]
    assert summaries[1]["distinctCount"] == 2
    assert summaries[1]["topValues"] == [
        {"value": '{"x":1}', "count": 2, "selectionValue": None},
        {"value": '{"x":2}', "count": 1, "selectionValue": None},
    ]


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_column_values_excludes_null_and_nan_special_values(lazy: bool):
    frame = pl.DataFrame({"value": [1.0, None, float("nan"), float("nan")]})
    source = frame.lazy() if lazy else frame

    values, has_more = PolarsEngine().column_values(source, "value")

    assert values == [{"value": "1.0", "count": 1, "selectionValue": typed_selection_value(1.0, "float")}]
    assert has_more is False


def _polars_formula_literal_operation(
    frame: Any, operator: str, value: Any, *, right_column: bool = False
) -> dict[str, Any]:
    schema = PolarsEngine().schema(frame)
    lineage = source_lineage(schema)
    return bind_step(
        validate_step(
            {
                "id": "exact-literal",
                "kind": "formula",
                "params": {
                    "leftColumn": lineage[0],
                    "operator": operator,
                    **({"rightColumn": lineage[1]} if right_column else {"value": value}),
                    "newColumn": "result",
                },
            }
        ),
        schema,
        lineage,
    )


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize(
    ("left", "right", "operator", "value"),
    [
        (pl.Series([False, True, None]), pl.Series([2, 127, None], dtype=pl.Int8), "add", None),
        (pl.Series([2, 255, None], dtype=pl.UInt8), pl.Series([False, True, None]), "add", None),
        (pl.Series([True, False, None]), pl.Series([1, 2, None], dtype=pl.UInt8), "subtract", None),
        (pl.Series([2, -128, None], dtype=pl.Int8), pl.Series([False, True, None]), "subtract", None),
        (pl.Series([False, True, None]), None, "add", str(2**127 - 1)),
        (pl.Series([False, True, None]), None, "add", str(2**128 - 1)),
        (pl.Series([False, True, None]), None, "subtract", str(-(2**127))),
    ],
)
def test_polars_formula_boolean_operands_refuse_native_integer_overflow(
    lazy: bool, left: Any, right: Any, operator: str, value: str | None
) -> None:
    source = pl.DataFrame({"left": left, **({"right": right} if right is not None else {})})
    before = source.serialize()
    frame = source.lazy() if lazy else source
    engine = PolarsEngine()
    operation = _polars_formula_literal_operation(frame, operator, value, right_column=right is not None)
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([operation]), namespace)
    for run in (lambda: engine.apply_transform(frame, operation), lambda: namespace["clean_data"](frame)):
        with pytest.raises((EngineError, ValueError), match="native integer capacity"):
            run()
        assert source.serialize() == before


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize(
    ("left", "right", "operator", "value", "expected"),
    [
        (
            pl.Series([False, True, None]),
            pl.Series([-(2**63)] * 3, dtype=pl.Int64),
            "multiply",
            None,
            [0, -(2**63), None],
        ),
        (
            pl.Series([-(2**63)] * 3, dtype=pl.Int64),
            pl.Series([False, True, None]),
            "multiply",
            None,
            [0, -(2**63), None],
        ),
        (pl.Series([], dtype=pl.Boolean), pl.Series([], dtype=pl.Int8), "multiply", None, []),
        (pl.Series([None], dtype=pl.Boolean), pl.Series([255], dtype=pl.UInt8), "add", None, [None]),
        (pl.Series([False]), None, "add", str(2**127 - 1), [2**127 - 1]),
        (pl.Series([None], dtype=pl.Boolean), None, "add", str(2**128 - 1), [None]),
        (pl.Series([], dtype=pl.Boolean), None, "subtract", str(-(2**127)), []),
        (pl.Series([True, False, None]), pl.Series([False, True, None]), "add", None, [1, 1, None]),
        (pl.Series([True, False, None]), pl.Series([0.5, 1.5, None]), "add", None, [1.5, 1.5, None]),
        (
            pl.Series([Decimal("0.50"), Decimal("1.50"), None], dtype=pl.Decimal(10, 2)),
            None,
            "add",
            "2",
            [Decimal("2.50"), Decimal("3.50"), None],
        ),
        (pl.Series([True, False, None]), None, "divide", "2", [0.5, 0.0, None]),
        (pl.Series([True, False, None]), None, "modulo", "2", [1, 0, None]),
    ],
)
def test_polars_formula_boolean_operands_preserve_native_results(
    lazy: bool, left: Any, right: Any, operator: str, value: str | None, expected: list[Any]
) -> None:
    source = pl.DataFrame({"left": left, **({"right": right} if right is not None else {})})
    before = source.serialize()
    frame = source.lazy() if lazy else source
    engine = PolarsEngine()
    operation = _polars_formula_literal_operation(frame, operator, value, right_column=right is not None)
    if right is not None:
        operand = pl.col("right")
    else:
        assert value is not None
        integer = int(value)
        literal_type = pl.Int64 if -(2**63) <= integer < 2**63 else pl.Int128 if integer < 2**127 else pl.UInt128
        operand = pl.lit(value).cast(literal_type)
    native = source.with_columns(polars_engine._polars_formula(pl.col("left"), operand, operator).alias("result"))
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([operation]), namespace)
    for result in (engine.apply_transform(frame, operation), namespace["clean_data"](frame)):
        assert isinstance(result, pl.LazyFrame) == lazy
        eager = result.collect() if lazy else result
        assert eager["result"].to_list() == expected
        assert eager.schema == native.schema
        assert eager.equals(native)
        assert source.serialize() == before


@pytest.mark.parametrize("boolean_base", [False, True])
def test_polars_formula_boolean_power_keeps_native_refusal(boolean_base: bool) -> None:
    columns = {"boolean": [True, False, None], "integer": [2, 3, None]}
    source = pl.DataFrame(columns if boolean_base else dict(reversed(list(columns.items()))))
    engine = PolarsEngine()
    operation = _polars_formula_literal_operation(source, "power", None, right_column=True)
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([operation]), namespace)
    for run in (lambda: engine.apply_transform(source, operation), lambda: namespace["clean_data"](source)):
        with pytest.raises(pl.exceptions.InvalidOperationError, match="not supported.*bool"):
            run()


@pytest.mark.parametrize(
    ("right", "operator", "error", "message"),
    [
        (pl.Series([False, True, None]), "multiply", pl.exceptions.InvalidOperationError, "not supported.*bool"),
        (
            pl.Series([Decimal("0.50"), Decimal("1.50"), None]),
            "add",
            pl.exceptions.SchemaError,
            "supertype of bool and decimal",
        ),
    ],
)
def test_polars_formula_boolean_unsupported_types_keep_native_refusal(
    right: Any, operator: str, error: type[Exception], message: str
) -> None:
    source = pl.DataFrame({"left": [True, False, None], "right": right})
    before = source.serialize()
    engine = PolarsEngine()
    operation = _polars_formula_literal_operation(source, operator, None, right_column=True)
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([operation]), namespace)
    for run in (lambda: engine.apply_transform(source, operation), lambda: namespace["clean_data"](source)):
        with pytest.raises(error, match=message):
            run()
        assert source.serialize() == before


def test_polars_saved_formula_refuses_boolean_source_replay_without_publishing(tmp_path: Path) -> None:
    path = tmp_path / "saved-formula.csv"
    path.write_text("a,b\n-1,-9223372036854775808\n", encoding="utf-8")
    manager = SessionManager()
    try:
        opened = manager.open_session({"kind": "file", "path": str(path)}, backend="polars", page_size=1)
        session_id = opened["metadata"]["sessionId"]
        left, right = opened["metadata"]["schema"]
        step = {
            "id": "saved-formula",
            "kind": "formula",
            "params": {
                "leftColumn": {"id": left["id"], "name": left["name"]},
                "rightColumn": {"id": right["id"], "name": right["name"]},
                "operator": "subtract",
                "newColumn": "result",
            },
        }
        manager.preview_step(session_id, 0, step, 0, 1)
        confirmed = manager.apply_draft(session_id, 1, 0, 1)
        assert confirmed["page"]["rows"][0]["values"][-1]["raw"] == str(2**63 - 1)
        namespace: dict[str, Any] = {}
        exec(confirmed["code"], namespace)
        assert namespace["clean_data"](pl.read_csv(path))["result"].to_list() == [2**63 - 1]
        saved_plan = deepcopy(manager.sessions[session_id].plan)
        manager.close_session(session_id, 2)

        source_bytes = b"a,b\ntrue,-9223372036854775808\n"
        path.write_bytes(source_bytes)
        reopened = manager.open_session({"kind": "file", "path": str(path)}, backend="polars", page_size=1)
        session_id = reopened["metadata"]["sessionId"]
        assert [column["type"] for column in reopened["metadata"]["schema"]] == ["boolean", "integer"]
        session = manager.sessions[session_id]
        committed = session.committed
        with pytest.raises(EngineError, match="native integer capacity"):
            manager.preview_step(session_id, 0, saved_plan[0], 0, 1)
        with pytest.raises(ValueError, match="native integer capacity"):
            namespace["clean_data"](pl.read_csv(path))
        assert session.revision == 0
        assert session.committed is committed
        assert session.plan == [] and session.bound_plan == []
        assert session.draft_step is None and session.draft_frame is None
        current = manager.get_page(session_id, 0, 0, 1, {"filters": [], "sort": []})
        assert current["metadata"] == reopened["metadata"]
        assert current["page"] == reopened["page"]
        assert saved_plan == [step]
        assert path.read_bytes() == source_bytes
    finally:
        manager.close_all()


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize(
    ("dtype", "values", "value", "operator", "expected", "result_dtype"),
    [
        (pl.Int64, [0, 7, None], 2**63, "add", [2**63, 2**63 + 7, None], pl.UInt64),
        (pl.Int64, [0, 7, None], 2**126, "add", [2**126, 2**126 + 7, None], pl.Int128),
        (pl.Int128, [0, 7, None], 2**127 - 8, "add", [2**127 - 8, 2**127 - 1, None], pl.Int128),
        (pl.Int128, [-1, 1, None], 2**127 - 1, "add", None, None),
        (pl.Int128, [1, None], 2**127 - 1, "add", [2**127, None], pl.UInt128),
        (pl.UInt128, [0, None], 2**128 - 1, "add", [2**128 - 1, None], pl.UInt128),
        (pl.UInt128, [1, None], 2**128 - 1, "add", None, None),
        (pl.Int128, [-1, None], 2**127, "add", None, None),
        (pl.Int8, [120, None], 10, "add", [130, None], pl.UInt8),
        (pl.UInt8, [0, 255, None], 1, "subtract", [-1, 254, None], pl.Int16),
        (pl.UInt8, [1, 255, None], -2, "multiply", [-2, -510, None], pl.Int16),
        (pl.Int128, [2**126, None], 4, "multiply", None, None),
        (pl.Int128, [-(2**127), None], -1, "modulo", [0, None], pl.Int128),
        (pl.Int64, [-7, 7, None], -3, "modulo", [-1, -2, None], pl.Int64),
        (pl.Int64, [7, None], 0, "modulo", [None, None], pl.Int64),
        (pl.Int8, [-2, None], 127, "power", [-(2**127), None], pl.Int128),
        (pl.Int8, [2, None], 127, "power", [2**127, None], pl.UInt128),
        (pl.Int8, [2, None], 128, "power", None, None),
        (pl.Int8, [-1, 0, 1, None], 2**32 - 1, "power", [-1, 0, 1, None], pl.Int64),
        (pl.Int8, [1, None], 2**32, "power", None, None),
        (pl.Int8, [2, None], -1, "power", None, None),
        (pl.UInt128, [None], -1, "add", [None], pl.Int128),
        (pl.Int64, [], 2**126, "add", [], pl.Int128),
        (pl.Int8, [None], 2**32, "power", None, None),
    ],
)
def test_polars_formula_integer_string_native_capacity(
    lazy: bool,
    dtype: Any,
    values: list[Any],
    value: int,
    operator: str,
    expected: list[Any] | None,
    result_dtype: Any,
) -> None:
    source = pl.DataFrame({"minimum": pl.Series(values, dtype=dtype)})
    before = source.clone()
    frame = source.lazy() if lazy else source
    engine = PolarsEngine()
    operation = _polars_formula_literal_operation(frame, operator, str(value))
    namespace: dict[str, Any] = {}
    code = engine.compile_plan([operation])
    exec(compile(code, "<generated>", "exec", dont_inherit=True), namespace)
    for run, error_type in (
        (lambda: engine.apply_transform(frame, operation), EngineError),
        (lambda: namespace["clean_data"](frame), ValueError),
    ):
        if expected is None:
            with pytest.raises(error_type, match="native integer.*capacity") as error:
                run()
            assert type(error.value) is error_type
        else:
            result = run()
            assert isinstance(result, pl.LazyFrame) == lazy
            eager = result.collect() if lazy else result
            assert eager["result"].to_list() == expected
            assert eager.schema["result"] == result_dtype
            assert eager.select(source.columns).equals(before)
        assert source.equals(before)
    assert operation["params"]["value"] == str(value)


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize("dtype", [pl.Int64, pl.Float64])
def test_polars_formula_integer_string_division_retains_native_float(lazy: bool, dtype: Any) -> None:
    source = pl.DataFrame({"value": pl.Series([0, 7, None], dtype=dtype)})
    frame = source.lazy() if lazy else source
    engine = PolarsEngine()
    value = 2**63 + 1
    operation = _polars_formula_literal_operation(frame, "divide", str(value))
    expected = source.with_columns((pl.col("value") / pl.lit(str(value)).cast(pl.Int128)).alias("result"))
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([operation]), namespace)
    for result in (engine.apply_transform(frame, operation), namespace["clean_data"](frame)):
        assert isinstance(result, pl.LazyFrame) == lazy
        assert (result.collect() if lazy else result).equals(expected)
        assert (result.collect_schema() if lazy else result.schema)["result"] == pl.Float64


@pytest.mark.parametrize("value", [str(2**128), str(-(2**127) - 1), "1" + "0" * 308, "1\n"])
def test_polars_formula_integer_string_refuses_invalid_literal(value: str) -> None:
    engine = PolarsEngine()
    frame = pl.DataFrame({"value": [1, None]}).lazy()
    operation = _polars_formula_literal_operation(frame, "add", "1")
    operation["params"]["value"] = value
    if value == "1\n":
        for run in (lambda: engine.apply_transform(frame, operation), lambda: engine.compile_plan([operation])):
            with pytest.raises(OperationError, match="Formula integer text must be canonical") as error:
                run()
            assert type(error.value) is OperationError
        return
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([operation]), namespace)
    for run, error_type in (
        (lambda: engine.apply_transform(frame, operation), EngineError),
        (lambda: namespace["clean_data"](frame), ValueError),
    ):
        with pytest.raises(error_type, match="native integer.*capacity") as error:
            run()
        assert type(error.value) is error_type


@pytest.mark.parametrize("right_column", [False, True])
def test_polars_formula_native_integer_collects_only_one_guard_boolean(
    right_column: bool, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = pl.DataFrame({"value": pl.Series([100, None], dtype=pl.Int8), "other": pl.Series([10, 2], dtype=pl.Int8)})
    frame = source.lazy()
    engine = PolarsEngine()
    operation = _polars_formula_literal_operation(frame, "add", 10)
    right = pl.lit(10)
    if right_column:
        schema = engine.schema(frame)
        lineage = source_lineage(schema)
        operation = bind_step(
            validate_step(
                {
                    "id": "column-formula",
                    "kind": "formula",
                    "params": {
                        "leftColumn": lineage[0],
                        "rightColumn": lineage[1],
                        "operator": "add",
                        "newColumn": "result",
                    },
                }
            ),
            schema,
            lineage,
        )
        right = pl.col("other")
    expected = source.with_columns((pl.col("value") + right).alias("result"))
    namespace: dict[str, Any] = {}
    code = engine.compile_plan([operation])
    exec(compile(code, "<generated>", "exec", dont_inherit=True), namespace)
    native_collect = pl.LazyFrame.collect
    observed: list[tuple[dict[str, Any], list[tuple[Any, ...]]]] = []

    def guard_collect(query: Any, *args: Any, **kwargs: Any) -> Any:
        result = cast(pl.DataFrame, native_collect(query, *args, **kwargs))
        observed.append((dict(result.schema), result.rows()))
        return result

    with monkeypatch.context() as guard:
        guard.setattr(pl, "collect_all", lambda *_args, **_kwargs: pytest.fail("Formula collected extra queries."))
        guard.setattr(pl.LazyFrame, "collect", guard_collect)
        live = engine.apply_transform(frame, operation)
        assert observed == [({"invalid": pl.Boolean}, [(False,)])]
        observed.clear()
        generated = namespace["clean_data"](frame)
    assert observed == [({"invalid": pl.Boolean}, [(False,)])]
    assert live.collect().equals(expected)
    assert generated.collect().equals(expected)
    assert live.collect_schema() == generated.collect_schema() == expected.schema


def test_polars_formula_integer_string_preflights_only_selected_native_bounds(monkeypatch: pytest.MonkeyPatch) -> None:
    source = pl.DataFrame({"minimum": list(range(100)), "maximum": ["untouched"] * 100})
    frame = source.lazy()
    engine = PolarsEngine()
    operation = _polars_formula_literal_operation(frame, "add", str(2**63))
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([operation]), namespace)
    native_collect = pl.LazyFrame.collect
    observed: list[tuple[dict[str, Any], list[tuple[Any, ...]]]] = []

    def bounded_collect(query: Any, *args: Any, **kwargs: Any) -> Any:
        result = cast(pl.DataFrame, native_collect(query, *args, **kwargs))
        observed.append((dict(result.schema), result.rows()))
        return result

    with monkeypatch.context() as guard:
        guard.setattr(pl, "collect_all", lambda *_args, **_kwargs: pytest.fail("Formula collected extra queries."))
        guard.setattr(pl.LazyFrame, "collect", bounded_collect)
        live = engine.apply_transform(frame, operation)
        assert observed == [({"minimum": pl.Int64, "maximum": pl.Int64}, [(0, 99)])]
        observed.clear()
        generated = namespace["clean_data"](frame)
    assert observed == [({"minimum": pl.Int64, "maximum": pl.Int64}, [(0, 99)])]
    for result in (live, generated):
        eager = result.collect()
        assert eager["result"].to_list() == [2**63 + value for value in range(100)]
        assert eager.select(source.columns).equals(source)


def test_polars_formula_integer_string_small_page_overflow_preserves_confirmed_session(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import __main__

    source = pl.DataFrame({"value": pl.Series([-1, 2**127 - 1], dtype=pl.Int128)})
    before = source.clone()
    monkeypatch.setattr(__main__, "formula_bounds_source", source.lazy(), raising=False)
    manager = SessionManager()
    opened = manager.open_session(
        {"kind": "notebookVariable", "label": "formula_bounds_source", "variableName": "formula_bounds_source"},
        backend="polars",
        mode="editing",
        page_size=1,
    )
    session_id = opened["metadata"]["sessionId"]
    session = manager.sessions[session_id]
    committed = session.committed
    column = opened["metadata"]["schema"][0]
    try:
        with pytest.raises(EngineError, match="native integer.*capacity"):
            manager.preview_step(
                session_id,
                0,
                {
                    "id": "overflow",
                    "kind": "formula",
                    "params": {
                        "leftColumn": {"id": column["id"], "name": column["name"]},
                        "operator": "add",
                        "value": "1",
                        "newColumn": "result",
                    },
                },
                0,
                1,
            )
        assert session.revision == 0
        assert session.committed is committed
        assert session.draft_step is None
        assert session.draft_frame is None
        assert session.plan == []
        assert session.bound_plan == []
        current = manager.get_page(session_id, 0, 0, 1, {"filters": [], "sort": []})
        assert current["metadata"] == opened["metadata"]
        assert current["page"] == opened["page"]
        assert source.equals(before)
    finally:
        manager.close_all()


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize(
    ("dtype", "values", "first_value", "first_dtype", "second_value", "expected", "final_dtype"),
    [
        (pl.Int8, [120, None], "10", pl.UInt8, "200", [-70, None], pl.Int16),
        (pl.Int64, [0, 7, None], str(2**63), pl.UInt64, str(2**63 + 1), [-1, 6, None], pl.Int128),
    ],
)
def test_polars_formula_integer_string_downstream_signedness(
    lazy: bool,
    dtype: Any,
    values: list[Any],
    first_value: str,
    first_dtype: Any,
    second_value: str,
    expected: list[Any],
    final_dtype: Any,
) -> None:
    source = pl.DataFrame({"value": pl.Series(values, dtype=dtype)})
    before = source.clone()
    frame = source.lazy() if lazy else source
    engine = PolarsEngine()
    first_step = _polars_formula_literal_operation(frame, "add", first_value)
    intermediate = engine.apply_transform(frame, first_step)
    assert (intermediate.collect_schema() if lazy else intermediate.schema)["result"] == first_dtype
    schema = engine.schema(intermediate)
    lineage = source_lineage(schema)
    second_step = bind_step(
        validate_step(
            {
                "id": "negative-result",
                "kind": "formula",
                "params": {
                    "leftColumn": lineage[1],
                    "value": second_value,
                    "operator": "subtract",
                    "newColumn": "next",
                },
            }
        ),
        schema,
        lineage,
    )
    live = engine.apply_transform(intermediate, second_step)
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([first_step, second_step]), namespace)
    generated = namespace["clean_data"](frame)
    for result in (live, generated):
        assert isinstance(result, pl.LazyFrame) == lazy
        eager = result.collect() if lazy else result
        assert eager["next"].to_list() == expected
        assert eager.schema["next"] == final_dtype
        assert eager.select(source.columns).equals(before)
    assert source.equals(before)


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize(
    ("left_dtype", "left_values", "right_dtype", "right_values", "operator"),
    [
        (pl.Int8, [1, 120, None], None, 10, "add"),
        (pl.UInt8, [2, 0, None], None, 1, "subtract"),
        (pl.Int16, [2, 30000, None], None, 3, "multiply"),
        (pl.Int8, [1, None], None, 2**127 - 1, "add"),
        (pl.Int128, [0, 2**127 - 1, None], pl.Int128, [1, 1, None], "add"),
        (pl.Int128, [0, -(2**127), None], pl.Int128, [1, 1, None], "subtract"),
        (pl.Int128, [1, 2**127 - 1, None], pl.Int128, [2, 2, None], "multiply"),
        (pl.Int128, [2, 2**64, None], pl.UInt32, [3, 2, None], "power"),
        (pl.UInt128, [1, 2, None], pl.UInt128, [2**32 - 1, 128, None], "power"),
        (pl.Int8, [2, None], None, 7, "power"),
        (pl.Int128, [None], pl.UInt128, [2**32], "power"),
        (pl.Int128, [-1, None], pl.UInt128, [2**127, None], "add"),
        (pl.UInt64, [2**64 - 1, None], pl.Int64, [-(2**63), None], "add"),
        (pl.UInt64, [2**64 - 1, None], pl.Int64, [2, None], "modulo"),
    ],
)
def test_polars_formula_native_integer_rejects_unsafe_rows(
    lazy: bool,
    left_dtype: Any,
    left_values: list[Any],
    right_dtype: Any,
    right_values: Any,
    operator: str,
) -> None:
    source = pl.DataFrame({"value": pl.Series(left_values, dtype=left_dtype)})
    if right_dtype is not None:
        source = source.with_columns(pl.Series("right", right_values, dtype=right_dtype))
    before = source.clone()
    frame = source.lazy() if lazy else source
    engine = PolarsEngine()
    operation = _polars_formula_literal_operation(frame, operator, right_values, right_column=right_dtype is not None)
    engine.validate_transform_preflight(frame, operation, engine.shape(frame))
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([operation]), namespace)
    for run, error_type in (
        (lambda: engine.apply_transform(frame, operation), EngineError),
        (lambda: namespace["clean_data"](frame), ValueError),
    ):
        with pytest.raises(error_type, match="native integer capacity|loses precision") as error:
            run()
        assert type(error.value) is error_type
        assert source.equals(before)


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize(
    ("left_dtype", "left_values", "right_dtype", "right_values", "operator", "expected", "dtype"),
    [
        (pl.Int8, [120, None], None, 128, "add", [248, None], pl.Int16),
        (pl.Int8, [120, None], None, 10.0, "add", [130.0, None], pl.Float64),
        (pl.Int8, [7, None], None, 2, "divide", [3.5, None], pl.Float64),
        (
            pl.Int128,
            [2**127 - 1, 0, None],
            pl.Int128,
            [-(2**127) + 1, 2**127 - 1, 1],
            "add",
            [0, 2**127 - 1, None],
            pl.Int128,
        ),
        (
            pl.Int128,
            [-(2**127), 2**127 - 1, None],
            pl.Int128,
            [-(2**127), 2**127 - 1, 1],
            "subtract",
            [0, 0, None],
            pl.Int128,
        ),
        (
            pl.Int128,
            [2**127 - 1, 1, 0, None],
            pl.Int128,
            [1, 2**127 - 1, 2**127 - 1, 2],
            "multiply",
            [2**127 - 1, 2**127 - 1, 0, None],
            pl.Int128,
        ),
        (pl.Int128, [-(2**127), None], None, 1, "multiply", [-(2**127), None], pl.Int128),
        (pl.UInt128, [2**128 - 1, None], None, 1, "multiply", [2**128 - 1, None], pl.UInt128),
        (
            pl.Int128,
            [-2, 2, 1, 0, None],
            pl.UInt32,
            [127, 3, 2**32 - 1, 1, 2],
            "power",
            [-(2**127), 8, 1, 0, None],
            pl.Int128,
        ),
        (pl.Int8, [-2, 2, 0, None], pl.Int64, [7, 6, 0, 1], "power", [-128, 64, 1, None], pl.Int8),
        (pl.Int8, [4, None], None, 0.5, "power", [2.0, None], pl.Float64),
        (pl.Int128, [None, -1], pl.UInt128, [2**127, None], "add", [None, None], pl.Int128),
        (pl.UInt64, [2**64 - 1, None], pl.Int64, [1, None], "add", [float(2**64), None], pl.Float64),
        (pl.UInt64, [9, 1, None], pl.Int64, [-4, 0, 2], "modulo", [-3.0, float("nan"), None], pl.Float64),
        (pl.Int128, [1, None], None, 0, "modulo", [None, None], pl.Int128),
        (pl.Int128, [], pl.UInt128, [], "multiply", [], pl.Int128),
        (pl.UInt128, [None, None], pl.UInt128, [0, None], "power", [None, None], pl.UInt128),
        (pl.Decimal(12, 2), [Decimal("1.25"), None], None, 2, "multiply", [Decimal("2.50"), None], pl.Decimal(38, 2)),
    ],
)
def test_polars_formula_native_integer_keeps_correlated_and_native_results(
    lazy: bool,
    left_dtype: Any,
    left_values: list[Any],
    right_dtype: Any,
    right_values: Any,
    operator: str,
    expected: list[Any],
    dtype: Any,
) -> None:
    source = pl.DataFrame({"value": pl.Series(left_values, dtype=left_dtype)})
    if right_dtype is not None:
        source = source.with_columns(pl.Series("right", right_values, dtype=right_dtype))
    before = source.clone()
    frame = source.lazy() if lazy else source
    engine = PolarsEngine()
    operation = _polars_formula_literal_operation(frame, operator, right_values, right_column=right_dtype is not None)
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([operation]), namespace)
    for result in (engine.apply_transform(frame, operation), namespace["clean_data"](frame)):
        assert isinstance(result, pl.LazyFrame) == lazy
        eager = result.collect(engine="streaming") if lazy else result
        assert eager["result"].equals(pl.Series("result", expected, dtype=dtype))
        assert eager.select(source.columns).equals(before)
        assert source.equals(before)


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_formula_native_integer_hidden_overflow_restores_session(
    monkeypatch: pytest.MonkeyPatch, lazy: bool
) -> None:
    import __main__

    source = pl.DataFrame({"value": pl.Series([1, 120], dtype=pl.Int8)})
    before = source.clone()
    monkeypatch.setattr(__main__, "formula_native_source", source.lazy() if lazy else source, raising=False)
    manager = SessionManager()
    opened = manager.open_session(
        {"kind": "notebookVariable", "variableName": "formula_native_source"},
        backend="polars",
        mode="editing",
        page_size=1,
    )
    session_id = opened["metadata"]["sessionId"]
    session = manager.sessions[session_id]
    committed = session.committed
    column = opened["metadata"]["schema"][0]
    try:
        with pytest.raises(EngineError, match="native integer capacity"):
            manager.preview_step(
                session_id,
                0,
                {
                    "id": "hidden-overflow",
                    "kind": "formula",
                    "params": {
                        "leftColumn": {"id": column["id"], "name": column["name"]},
                        "operator": "add",
                        "value": 10,
                        "newColumn": "result",
                    },
                },
                0,
                1,
            )
        assert session.revision == 0
        assert session.committed is committed
        assert session.draft_step is None and session.draft_frame is None
        assert session.plan == [] and session.bound_plan == []
        current = manager.get_page(session_id, 0, 0, 1, {"filters": [], "sort": []})
        assert current["metadata"] == opened["metadata"] and current["page"] == opened["page"]
        assert source.equals(before)
    finally:
        manager.close_all()


@pytest.mark.parametrize("lazy", [False, True])
def test_polars_formula_native_integer_mixed_plan_isolates_generated_helpers(lazy: bool) -> None:
    source = pl.DataFrame({"value": pl.Series([3, -4, None], dtype=pl.Int8), "invalid": [False, True, None]})
    before = source.clone()
    frame = source.lazy() if lazy else source
    engine = PolarsEngine()
    plan = [_polars_formula_literal_operation(frame, "add", 2)]
    current = engine.apply_transform(frame, plan[0])
    custom = validate_step(
        {
            "id": "helper-collision",
            "kind": "customCode",
            "params": {
                "code": (
                    "def _ow_polars_check_formula(*args):\n    raise AssertionError('custom local escaped')\n"
                    "def root_limit(*args):\n    return 0\n"
                    "result = df.with_columns(pl.col('result').alias('copied'))"
                )
            },
        }
    )
    schema = engine.schema(current)
    plan.append(bind_step(custom, schema, source_lineage(schema)))
    current = engine.apply_transform(current, plan[-1])
    for value, operator, output in [("200", "add", "wide"), (3, "power", "cube")]:
        schema = engine.schema(current)
        lineage = source_lineage(schema)
        selected = next(reference for reference in lineage if reference["name"] == "copied")
        step = bind_step(
            validate_step(
                {
                    "id": output,
                    "kind": "formula",
                    "params": {
                        "leftColumn": selected,
                        "value": value,
                        "operator": operator,
                        "newColumn": output,
                    },
                }
            ),
            schema,
            lineage,
        )
        plan.append(step)
        current = engine.apply_transform(current, step)
    namespace: dict[str, Any] = {}
    code = engine.compile_plan(plan)
    exec(code, namespace)
    generated = namespace["clean_data"](frame)
    for result in (current, generated):
        assert isinstance(result, pl.LazyFrame) == lazy
        eager = result.collect(engine="streaming") if lazy else result
        assert eager["cube"].equals(pl.Series("cube", [125, -8, None], dtype=pl.Int8))
        assert eager["wide"].equals(pl.Series("wide", [205, 198, None], dtype=pl.Int16))
        assert eager.select(source.columns).equals(before)
    assert source.equals(before)


@pytest.mark.parametrize("dtype", [pl.Float32, pl.Float64, pl.Decimal(12, 2)])
def test_polars_formula_native_noninteger_has_no_row_guard(dtype: Any, monkeypatch: pytest.MonkeyPatch) -> None:
    source = pl.DataFrame({"value": pl.Series([1, 2, None], dtype=dtype)})
    frame = source.lazy()
    engine = PolarsEngine()
    step = _polars_formula_literal_operation(frame, "multiply", 3)
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([step]), namespace)

    def reject_collect(*_args: Any, **_kwargs: Any) -> Any:
        pytest.fail("Floating/Decimal rows scanned during Formula construction.")

    with monkeypatch.context() as guard:
        guard.setattr(pl, "collect_all", reject_collect)
        guard.setattr(pl.LazyFrame, "collect", reject_collect)
        live = engine.apply_transform(frame, step)
        generated = namespace["clean_data"](frame)
    expected = source.with_columns((pl.col("value") * pl.lit(3)).alias("result"))
    assert live.collect().equals(expected)
    assert generated.collect().equals(expected)
    assert live.collect_schema() == generated.collect_schema() == expected.schema


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize("operator", ["add", "subtract", "multiply"])
def test_polars_uint128_column_formula_has_correlated_preview_and_safe_later_pages(
    lazy: bool, operator: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    import __main__
    from openwrangler_runtime import kernel_agent

    source = pl.DataFrame(
        {"value": pl.Series([1, 2, None], dtype=pl.UInt64), "right": pl.Series([3, 4, None], dtype=pl.UInt64)}
    )
    original = source.clone()
    frame = source.lazy() if lazy else source
    monkeypatch.setattr(__main__, "uint128_formula_source", frame, raising=False)
    manager = SessionManager()
    monkeypatch.setattr(kernel_agent, "_manager", manager)
    opened = manager.open_session(
        {"kind": "notebookVariable", "variableName": "uint128_formula_source"},
        backend="polars",
        mode="editing",
        page_size=1,
    )
    session_id = opened["metadata"]["sessionId"]
    first = opened["metadata"]["schema"][0]
    try:
        preview = manager.preview_step(
            session_id,
            0,
            {
                "id": "wide",
                "kind": "formula",
                "params": {
                    "leftColumn": {"id": first["id"], "name": first["name"]},
                    "value": str(2**64),
                    "operator": "add",
                    "newColumn": "wide",
                },
            },
            0,
            1,
        )
        confirmed = manager.apply_draft(session_id, preview["revision"], 0, 3)
        session = manager.sessions[session_id]
        committed, revision, plan = session.committed, session.revision, deepcopy(session.plan)
        refs = {
            column["name"]: {"id": column["id"], "name": column["name"]} for column in confirmed["metadata"]["schema"]
        }
        operation = {
            "id": "column-arithmetic",
            "kind": "formula",
            "params": {
                "leftColumn": refs["wide"],
                "rightColumn": refs["right"],
                "operator": operator,
                "newColumn": "result",
            },
        }
        bound = bind_step(validate_step(operation), session.committed_schema, session.committed_lineage)
        namespace: dict[str, Any] = {}
        exec(session.engine.compile_plan([*session.bound_plan, bound]), namespace)
        request_id = f"uint128-{operator}-{lazy}"
        response = json.loads(
            kernel_agent.dispatch_json(
                json.dumps(
                    {
                        "protocolVersion": 4,
                        "requestId": request_id,
                        "priority": "interactive",
                        "request": {
                            "kind": "previewStep",
                            "sessionId": session_id,
                            "revision": revision,
                            "step": operation,
                            "offset": 0,
                            "limit": 1,
                            "columnOffset": 0,
                            "columnLimit": 64,
                        },
                    }
                )
            )
        )
        assert response["requestId"] == request_id
        if pl.__version__.startswith("1.35."):
            assert response["response"]["kind"] == "error"
            assert "Polars 1.36" in response["response"]["message"]
            assert session.committed is committed and session.revision == revision and session.plan == plan
            assert session.draft_step is None and session.draft_frame is None
            actual = manager.get_page(session_id, revision, 0, 3, {"filters": [], "sort": []})
            assert actual["metadata"] == confirmed["metadata"] and actual["page"] == confirmed["page"]
            with pytest.raises(ValueError, match="Polars 1.36"):
                namespace["clean_data"](frame)
            # A corrected scalar operation still works on the preserved UInt128 result.
            corrected = {**operation, "params": {**operation["params"], "value": 1}}
            corrected["params"].pop("rightColumn")
            retry = manager.preview_step(session_id, revision, corrected, 0, 1)
            manager.apply_draft(session_id, retry["revision"], 0, 3)
        else:
            assert response["response"]["kind"] == "stepPreview"
            manager.apply_draft(session_id, response["response"]["revision"], 0, 1)
            page = manager.get_page(session_id, session.revision, 0, 3, {"filters": [], "sort": []})
            generated = namespace["clean_data"](frame)
            generated = generated.collect() if lazy else generated
            expected = {
                "add": [2**64 + 4, 2**64 + 6, None],
                "subtract": [2**64 - 2, 2**64 - 2, None],
                "multiply": [(2**64 + 1) * 3, (2**64 + 2) * 4, None],
            }[operator]
            assert generated["result"].to_list() == expected and generated.schema["result"] == pl.UInt128
            assert [row["values"][-1]["raw"] for row in page["page"]["rows"]] == [
                str(value) if value is not None else None for value in expected
            ]
        assert source.equals(original)
    finally:
        manager.close_all()


@pytest.mark.parametrize("release", ["1.35.2", "1.36.0rc1", "custom"])
@pytest.mark.parametrize("values", [[], [None, None], [1], [1, 2]])
def test_polars_uint128_column_formula_refuses_unqualified_release_for_every_shape(
    release: str, values: list[int | None], monkeypatch: pytest.MonkeyPatch
) -> None:
    source = pl.DataFrame({"value": pl.Series(values, dtype=pl.UInt128), "right": pl.Series(values, dtype=pl.UInt128)})
    original = source.clone()
    engine = PolarsEngine()
    operation = _polars_formula_literal_operation(source, "add", None, right_column=True)
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([operation]), namespace)
    monkeypatch.setattr(pl, "__version__", release)
    for frame in (source, source.lazy()):
        for run in (
            lambda frame=frame: engine.apply_transform(frame, operation),
            lambda frame=frame: namespace["clean_data"](frame),
        ):
            with pytest.raises((EngineError, ValueError), match="stable Polars 1.36"):
                run()
    assert source.equals(original)


@pytest.mark.parametrize("other_dtype", [pl.Boolean, pl.Null])
@pytest.mark.parametrize("operator", ["add", "subtract", "multiply"])
@pytest.mark.parametrize("unsigned_left", [False, True])
def test_polars_uint128_kernel_refusal_includes_noninteger_operands(
    other_dtype: Any, operator: str, unsigned_left: bool, monkeypatch: pytest.MonkeyPatch
) -> None:
    unsigned = pl.Series([3, None], dtype=pl.UInt128)
    other = pl.Series([True, None] if other_dtype == pl.Boolean else [None, None], dtype=other_dtype)
    source = pl.DataFrame(
        {"value": unsigned if unsigned_left else other, "right": other if unsigned_left else unsigned}
    )
    original = source.clone()
    engine = PolarsEngine()
    operation = _polars_formula_literal_operation(source, operator, None, right_column=True)
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([operation]), namespace)
    monkeypatch.setattr(pl, "__version__", "1.35.2")
    for frame in (source, source.lazy()):
        for run in (
            lambda frame=frame: engine.apply_transform(frame, operation),
            lambda frame=frame: namespace["clean_data"](frame),
        ):
            with pytest.raises((EngineError, ValueError), match="stable Polars 1.36"):
                run()
    assert source.equals(original)


@pytest.mark.parametrize("boundary", ["live", "generated", "generated-drop"])
def test_polars_custom_result_refuses_invalid_expression_before_projection(boundary: str) -> None:
    source = pl.DataFrame({"pos": [0, 1, 2], "value": ["1", "bad", None]})
    before = source.clone()
    engine = PolarsEngine()
    custom = validate_step(
        {
            "id": "strict-cast",
            "kind": "customCode",
            "params": {"code": 'result = df.lazy().with_columns(pl.col("value").cast(pl.Int64))'},
        }
    )
    intermediate = source.lazy().with_columns(pl.col("value").cast(pl.Int64))
    assert isinstance(intermediate, pl.LazyFrame)
    assert intermediate.collect_schema() == {"pos": pl.Int64, "value": pl.Int64}
    if boundary == "live":
        with pytest.raises(pl.exceptions.InvalidOperationError):
            engine.apply_transform(source, custom)
    else:
        plan = [custom]
        if boundary == "generated-drop":
            schema = engine.schema(intermediate)
            plan.append(
                bind_step(
                    validate_step(
                        {
                            "id": "drop-invalid",
                            "kind": "dropColumns",
                            "params": {"columns": [source_lineage(schema)[1]]},
                        }
                    ),
                    schema,
                    source_lineage(schema),
                )
            )
        namespace: dict[str, Any] = {}
        exec(engine.compile_plan(plan), namespace)
        with pytest.raises(pl.exceptions.InvalidOperationError):
            namespace["clean_data"](source)
    assert source.schema == before.schema
    assert source.equals(before)


@pytest.mark.parametrize("boundary", ["notebook", "custom"])
def test_polars_retained_lazy_result_keeps_page_identity_and_session_ownership(
    boundary: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import __main__

    source = pl.DataFrame({"key": range(8), "bucket": [0, 1] * 4, "value": [key * 10 for key in range(8)]})
    before = source.clone()
    visits: list[int] = []

    def rotate(frame: pl.DataFrame) -> pl.DataFrame:
        shift = len(visits) % frame.height
        visits.append(frame.height)
        return pl.concat([frame.slice(shift), frame.head(shift)])

    lazy = source.lazy().map_batches(rotate, predicate_pushdown=False, projection_pushdown=False, slice_pushdown=False)
    monkeypatch.setattr(__main__, "retained_polars_source", lazy, raising=False)
    monkeypatch.setattr(__main__, "retained_polars_rotate", rotate, raising=False)
    path = tmp_path / "retained.parquet"
    source.write_parquet(path)
    contents, stat = path.read_bytes(), path.stat()
    manager = SessionManager()
    view: dict[str, Any] = {"filters": [], "sort": []}
    request = (
        {"kind": "notebookVariable", "variableName": "retained_polars_source"}
        if boundary == "notebook"
        else {"kind": "file", "path": str(path)}
    )
    try:
        if boundary == "notebook":
            bad = lazy.rename({"key": "__open_wrangler_internal_row_id_user"})
            monkeypatch.setattr(__main__, "retained_polars_source", bad)
            with pytest.raises(EngineError, match="private row-identity prefix"):
                manager.open_session(request, backend="polars", page_size=2)
            assert visits == [] and not manager.sessions
            bad = lazy.map_batches(
                lambda batch: batch.with_columns(
                    (pl.col("key") + 91).cast(pl.UInt32).alias("__open_wrangler_internal_row_id_user")
                ),
                validate_output_schema=False,
                predicate_pushdown=False,
                projection_pushdown=False,
                slice_pushdown=False,
            )
            assert bad.collect_schema() == source.schema
            monkeypatch.setattr(__main__, "retained_polars_source", bad)
            with pytest.raises(EngineError, match="private row-identity prefix"):
                manager.open_session(request, backend="polars", page_size=2)
            assert visits == [8] and not manager.sessions
            assert __main__.retained_polars_source is bad and source.equals(before) and source.schema == before.schema
            visits.clear()
            monkeypatch.setattr(__main__, "retained_polars_source", lazy)
        opened = manager.open_session(request, backend="polars", mode="editing", page_size=2, column_limit=1)
        sid = opened["metadata"]["sessionId"]
        session = manager.sessions[sid]
        if boundary == "custom":
            assert "Parquet SCAN" in session.original.explain() and visits == []
            preview = manager.preview_step(
                sid,
                0,
                {
                    "id": "rotate",
                    "kind": "customCode",
                    "params": {
                        "code": (
                            "import __main__\nresult = df.map_batches(__main__.retained_polars_rotate, "
                            "predicate_pushdown=False, projection_pushdown=False, slice_pushdown=False)"
                        )
                    },
                },
                0,
                2,
                column_limit=1,
            )
            draft = session.draft_frame
            opened = manager.apply_draft(sid, preview["revision"], 0, 2, column_limit=1)
            assert session.committed is draft
        assert visits == [8]
        assert isinstance(session.original, pl.LazyFrame) and isinstance(session.committed, pl.LazyFrame)
        revision = session.revision
        full = manager.get_page(sid, revision, 0, 8, view, column_limit=3)
        rows = full["page"]["rows"]
        assert [row["id"] for row in opened["page"]["rows"]] == [row["id"] for row in rows[:2]]
        assert [[cell["raw"] for cell in row["values"]] for row in rows] == [
            [key, key % 2, key * 10] for key in range(8)
        ]
        ids = {row["values"][0]["raw"]: row["id"] for row in rows}
        assert len(set(ids.values())) == 8
        for offset in (0, 4):
            page = manager.get_page(sid, revision, offset, 4, view, column_limit=2)
            assert [(row["values"][0]["raw"], row["id"]) for row in page["page"]["rows"]] == [
                (key, ids[key]) for key in range(offset, offset + 4)
            ]
            assert manager.get_page(sid, revision, offset, 4, view, column_limit=2)["page"] is page["page"]
        filtered = {
            "logic": "and",
            "filters": [{"column": "key", "type": "integer", "predicates": [{"operator": "gte", "value": 2}]}],
            "sort": [{"column": "bucket", "direction": "asc", "nulls": "last"}],
        }
        page = manager.get_page(sid, revision, 0, 8, filtered, column_limit=1)
        assert [(row["values"][0]["raw"], row["id"]) for row in page["page"]["rows"]] == [
            (key, ids[key]) for key in (2, 4, 6, 3, 5, 7)
        ]
        confirmed, cache = session.committed, list(session.page_cache.items())
        for code, error_type in (
            (
                'result = df.with_columns(pl.lit("bad").cast(pl.Int64).alias("broken"))',
                pl.exceptions.InvalidOperationError,
            ),
            (
                "lazy = df.lazy() if isinstance(df, pl.DataFrame) else df\n"
                "result = lazy.map_batches(lambda batch: batch.select([]), "
                "predicate_pushdown=False, projection_pushdown=False, slice_pushdown=False, "
                "validate_output_schema=False, streamable=False)",
                EngineError,
            ),
        ):
            invalid = validate_step({"id": "invalid", "kind": "customCode", "params": {"code": code}})
            with pytest.raises(error_type) as error:
                manager.preview_step(sid, revision, invalid, 0, 2, column_limit=1)
            assert session.revision == revision and session.committed is confirmed and session.draft_frame is None
            assert list(session.page_cache.items()) == cache and session.filter_model == filtered
            assert manager.get_page(sid, revision, 0, 8, filtered, column_limit=1)["page"] is page["page"]
            if error_type is EngineError:
                assert str(error.value) == "A transformation must leave at least one visible column."
                generated_namespace: dict[str, Any] = {}
                exec(
                    compile(session.engine.compile_plan([invalid]), "<generated>", "exec", dont_inherit=True),
                    generated_namespace,
                )
                with pytest.raises(ValueError, match="^A transformation must leave at least one visible column[.]$"):
                    generated_namespace["clean_data"](source.lazy())
        assert visits == [8]
        if boundary == "notebook":
            clone = manager.open_session(
                request,
                backend="polars",
                mode="editing",
                page_size=2,
                clone_from={"sessionId": sid, "revision": revision},
            )
            clone_id = clone["metadata"]["sessionId"]
            assert manager.sessions[clone_id].original is session.original
            manager.close_session(sid, revision)
            cloned = manager.get_page(clone_id, 0, 0, 8, view, column_limit=3)
            assert cloned["page"]["rows"] == rows and visits == [8]
        else:
            undone = manager.undo_step(sid, revision, 0, 2)
            assert session.committed is session.original and visits == [8]
            redone = manager.redo_step(sid, undone["revision"], 0, 8)
            assert visits == [8, 8]
            assert [row["values"][0]["raw"] for row in redone["page"]["rows"]] == [2, 4, 6, 3, 5, 7]
            replayed = manager.get_page(sid, redone["revision"], 0, 8, view)
            assert [row["values"][0]["raw"] for row in replayed["page"]["rows"]] == [1, 2, 3, 4, 5, 6, 7, 0]
            visits.clear()
            namespace: dict[str, Any] = {}
            exec(compile(redone["code"], "<generated>", "exec", dont_inherit=True), namespace)
            generated = namespace["clean_data"](source.lazy())
            assert isinstance(generated, pl.LazyFrame) and visits == [8]
            assert generated.collect().equals(source) and generated.select("value").collect().equals(
                source.select("value")
            )
            assert visits == [8]
    finally:
        manager.close_all()
    assert not manager.sessions and source.equals(before) and source.schema == before.schema
    assert __main__.retained_polars_source is lazy
    after = path.stat()
    assert path.read_bytes() == contents
    assert (after.st_dev, after.st_ino, after.st_size, after.st_mtime_ns) == (
        stat.st_dev,
        stat.st_ino,
        stat.st_size,
        stat.st_mtime_ns,
    )


@pytest.mark.parametrize("lazy", [False, True])
@pytest.mark.parametrize("empty", [False, True])
def test_polars_custom_result_check_preserves_native_objects_nulls_and_empty_frames(lazy: bool, empty: bool) -> None:
    token = object()
    source = pl.DataFrame(
        {
            "object": pl.Series([token, None], dtype=pl.Object),
            "number": pl.Series([-0.0, None], dtype=pl.Float64),
            "nested": pl.Series([{"items": [1, None]}, None], dtype=pl.Struct({"items": pl.List(pl.Int64)})),
            "category": pl.Series(["b", None], dtype=pl.Enum(["a", "b", "unused"])),
            "binary": pl.Series([b"\x00\xff", None], dtype=pl.Binary),
        }
    )
    if empty:
        source = source.head(0)
    frame = source.lazy() if lazy else source
    engine = PolarsEngine()
    engine.validate_transformation_result(frame)
    step = validate_step({"id": "identity", "kind": "customCode", "params": {"code": "result = df"}})
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([step]), namespace)
    for result in (
        engine.apply_transform(frame, step),
        namespace["clean_data"](frame),
        engine.capture_notebook_source(frame),
    ):
        assert isinstance(result, type(frame))
        assert (result is frame) is not lazy
        output = result.collect(engine="streaming") if isinstance(result, pl.LazyFrame) else result
        assert output.schema == source.schema
        assert output.shape == source.shape
        assert output.null_count().row(0) == ((0, 0, 0, 0, 0) if empty else (1, 1, 1, 1, 1))
        if not empty:
            assert output["object"][0] is token and source["object"][0] is token
            assert output["number"][0].hex() == source["number"][0].hex() == "-0x0.0p+0"
            assert output["nested"].to_list() == source["nested"].to_list() == [{"items": [1, None]}, None]
            assert output["category"].to_list() == source["category"].to_list() == ["b", None]
            assert output["binary"].to_list() == source["binary"].to_list() == [b"\x00\xff", None]
    with pytest.raises(EngineError, match="at least one visible column"):
        engine.validate_transformation_result(pl.DataFrame().lazy() if lazy else pl.DataFrame())


@pytest.mark.parametrize("with_custom", [False, True])
def test_polars_only_custom_steps_evaluate_rows_during_plan_construction(with_custom: bool) -> None:
    visits: list[int | None] = []

    def observe(value: int | None) -> int | None:
        visits.append(value)
        return value

    source = pl.DataFrame({"value": pl.Series([1, 2, None], dtype=pl.Int64)})
    frame = source.lazy().with_columns(pl.col("value").map_elements(observe, return_dtype=pl.Int64, skip_nulls=False))
    engine = PolarsEngine()
    first_schema = engine.schema(frame)
    first = bind_step(
        validate_step(
            {
                "id": "first",
                "kind": "renameColumn",
                "params": {"column": source_lineage(first_schema)[0], "newName": "next"},
            }
        ),
        first_schema,
        source_lineage(first_schema),
    )
    intermediate = engine.apply_transform(frame, first)
    second_schema = engine.schema(intermediate)
    second = bind_step(
        validate_step(
            {
                "id": "second",
                "kind": "renameColumn",
                "params": {"column": source_lineage(second_schema)[0], "newName": "value"},
            }
        ),
        second_schema,
        source_lineage(second_schema),
    )
    plan = [first]
    if with_custom:
        plan.append(validate_step({"id": "custom", "kind": "customCode", "params": {"code": "result = df"}}))
    plan.append(second)
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan(plan), namespace)
    empty_namespace: dict[str, Any] = {}
    exec(engine.compile_plan([]), empty_namespace)
    assert empty_namespace["clean_data"](frame) is frame
    assert visits == []
    engine.validate_transformation_result(intermediate)
    assert visits == []
    if with_custom:
        retained = engine.apply_transform(intermediate, plan[1])
        assert isinstance(retained, pl.LazyFrame) and retained is not intermediate
        assert Counter(visits) == Counter([1, 2, None])
        assert retained.collect().rename({"next": "value"}).equals(source)
        assert Counter(visits) == Counter([1, 2, None])
    visits.clear()
    generated = namespace["clean_data"](frame)
    assert isinstance(generated, pl.LazyFrame)
    assert Counter(visits) == Counter([1, 2, None] if with_custom else [])
    visits.clear()
    assert generated.collect(engine="streaming").equals(source)
    assert Counter(visits) == Counter([] if with_custom else [1, 2, None])
    assert source.schema == {"value": pl.Int64} and source.rows() == [(1,), (2,), (None,)]


@pytest.mark.parametrize("population", ["values", "null", "empty"])
def test_polars_format_datetime_preserves_native_lazy_error_timing(population: str) -> None:
    dates = [0, None] if population == "values" else [None, None]
    source = pl.DataFrame({"pos": [0, 1], "date": pl.Series(dates, dtype=pl.Date)})
    if population == "empty":
        source = source.head(0)
    before = source.clone()
    frame = source.lazy()
    engine = PolarsEngine()
    schema = engine.schema(frame)
    lineage = source_lineage(schema)
    operation = bind_step(
        validate_step(
            {
                "id": "format",
                "kind": "formatDatetime",
                "params": {"column": lineage[1], "format": "%Q", "newColumn": "formatted"},
            }
        ),
        schema,
        lineage,
    )
    live = engine.apply_transform(frame, operation)
    engine.validate_transformation_result(live)
    namespace: dict[str, Any] = {}
    exec(engine.compile_plan([operation]), namespace)
    generated = namespace["clean_data"](frame)
    for result in (live, generated):
        assert isinstance(result, pl.LazyFrame)
        assert result.collect_schema() == {"pos": pl.Int64, "date": pl.Date, "formatted": pl.String}
        assert result.select("pos").collect().equals(source.select("pos"))
        if population == "values":
            with pytest.raises(pl.exceptions.PolarsError):
                result.collect()
        else:
            assert result.collect().equals(source.with_columns(pl.lit(None, dtype=pl.String).alias("formatted")))
    output_schema = engine.schema(live)
    output_lineage = source_lineage(output_schema)
    drop = bind_step(
        validate_step({"id": "drop", "kind": "dropColumns", "params": {"columns": [output_lineage[2]]}}),
        output_schema,
        output_lineage,
    )
    exec(engine.compile_plan([operation, drop]), namespace)
    assert namespace["clean_data"](frame).collect().equals(source)
    assert source.equals(before)
